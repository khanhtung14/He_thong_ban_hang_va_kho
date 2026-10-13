import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_sample_excel():
    wb = Workbook()
    
    # Sheet 1: Dữ liệu mẫu chuẩn tiếng Việt (Hệ thống tự động map theo alias)
    ws = wb.active
    ws.title = "Danh sách sản phẩm"
    
    headers = [
        ("Mã SKU", "sku"),
        ("Tên sản phẩm", "name"),
        ("Danh mục", "category"),
        ("Giá bán", "sale_price"),
        ("Giá vốn", "cost_price"),
        ("Tồn kho", "stock_available"),
        ("Đơn vị tính", "unit")
    ]
    
    sample_products = [
        # Nước giải khát
        ("NGK-COCA-330", "Nước ngọt Coca-Cola lon 330ml", "Nước giải khát", 10000, 7500, 250, "Lon"),
        ("NGK-PEPSI-330", "Nước ngọt Pepsi lon 330ml", "Nước giải khát", 10000, 7400, 180, "Lon"),
        ("NGK-REVIVE-500", "Nước điện giải Revive chanh muối 500ml", "Nước giải khát", 12000, 8800, 120, "Chai"),
        ("NGK-AQUAFINA-500", "Nước tinh khiết Aquafina 500ml", "Nước giải khát", 6000, 3800, 500, "Chai"),
        ("NGK-REDBULL-250", "Nước tăng lực Red Bull lon 250ml", "Nước giải khát", 14000, 10500, 300, "Lon"),
        
        # Bánh kẹo & Snack
        ("BK-COSY-ME-288", "Bánh quy mè Cosy Wonderfull 288g", "Bánh kẹo", 36000, 27000, 85, "Hộp"),
        ("BK-CHOCO-PIE-12", "Bánh ChocoPie hộp 12 cái (396g)", "Bánh kẹo", 58000, 44000, 110, "Hộp"),
        ("BK-KNOOP-OISHI", "Snack phồng tôm Oishi 40g", "Bánh kẹo", 7000, 4800, 200, "Gói"),
        ("BK-LAY-NATURAL", "Khoai tây chiên Lay's tự nhiên 95g", "Bánh kẹo", 22000, 16500, 95, "Gói"),
        
        # Sữa & Sản phẩm từ sữa
        ("SUA-VINAMILK-100", "Sữa tươi tiệt trùng Vinamilk có đường 100ml", "Sữa & Chế phẩm sữa", 7500, 5600, 400, "Hộp"),
        ("SUA-TH-TRUE-1L", "Sữa tươi thanh trùng TH True Milk nguyên chất 1L", "Sữa & Chế phẩm sữa", 38000, 29500, 70, "Hộp"),
        ("SUA-YAKULT-65", "Sữa chua uống lên men Yakult lốc 5 chai", "Sữa & Chế phẩm sữa", 26000, 20500, 150, "Lốc"),
        
        # Gia vị & Thực phẩm khô
        ("GV-CHINSU-500", "Nước mắm Chinsu đệ nhị 500ml", "Gia vị", 22000, 16800, 160, "Chai"),
        ("GV-MAGGI-300", "Dầu hào Maggi đậm đặc 300g", "Gia vị", 24000, 18000, 80, "Chai"),
        ("GV-AJINOMOTO-454", "Bột ngọt Ajinomoto hạt lớn gói 454g", "Gia vị", 35000, 28000, 140, "Gói"),
        ("GV-NEPTUNE-1L", "Dầu ăn cao cấp Neptune Light 1L", "Gia vị", 52000, 41500, 90, "Chai"),
        ("MI-HAOHAO-TMC", "Mì ăn liền Hảo Hảo tôm chua cay 75g", "Thực phẩm khô", 4500, 3200, 1200, "Gói"),
        ("MI-OMACHI-XB", "Mì khoai tây Omachi xốt bò hầm 80g", "Thực phẩm khô", 9000, 6800, 600, "Gói"),
        
        # Đồ dùng gia đình / Hóa phẩm
        ("HP-SUNLIGHT-750", "Nước rửa chén Sunlight chanh 750g", "Hóa phẩm", 31000, 23500, 130, "Chai"),
        ("HP-OMO-COMFORT-28", "Nước giặt OMO Matic Comfort tinh dầu 2.8kg", "Hóa phẩm", 185000, 145000, 45, "Túi"),
    ]
    
    # Styles
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
    alignment_center = Alignment(horizontal="center", vertical="center")
    alignment_left = Alignment(horizontal="left", vertical="center")
    alignment_right = Alignment(horizontal="right", vertical="center")
    
    thin_border_side = Side(border_style="thin", color="D1D5DB")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    font_body = Font(name="Segoe UI", size=10)
    
    # Header row
    for col_idx, (h_title, _) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_title)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = alignment_center
        cell.border = border_cell
    
    ws.row_dimensions[1].height = 28
    
    # Data rows
    for row_idx, prod in enumerate(sample_products, start=2):
        ws.row_dimensions[row_idx].height = 22
        is_even = (row_idx % 2 == 0)
        
        for col_idx, val in enumerate(prod, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_body
            cell.border = border_cell
            if is_even:
                cell.fill = fill_zebra
                
            # Formatting
            if col_idx in (1, 7): # SKU, Đơn vị tính
                cell.alignment = alignment_center
            elif col_idx in (2, 3): # Tên, Danh mục
                cell.alignment = alignment_left
            elif col_idx in (4, 5): # Giá bán, Giá vốn
                cell.alignment = alignment_right
                cell.number_format = "#,##0"
            elif col_idx == 6: # Tồn kho
                cell.alignment = alignment_right
                cell.number_format = "#,##0"

    # Auto fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)
    ws.column_dimensions['B'].width = 38 # Tên sản phẩm rộng hơn
    
    # Sheet 2: Bản tiếng Anh tương thích chuẩn API (Raw Keys)
    ws2 = wb.create_sheet(title="Template_Keys_EN")
    en_headers = ["sku", "name", "category", "sale_price", "cost_price", "stock_available", "unit"]
    for col_idx, h in enumerate(en_headers, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
        cell.alignment = alignment_center
        cell.border = border_cell
    
    for row_idx, prod in enumerate(sample_products[:5], start=2):
        for col_idx, val in enumerate(prod, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_body
            cell.border = border_cell
            if col_idx in (4, 5, 6):
                cell.number_format = "#,##0"

    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = 22
    ws2.column_dimensions['B'].width = 38
    
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    output_path = os.path.join(root_dir, "danh_sach_san_pham_mau.xlsx")
    wb.save(output_path)
    print(f"File created successfully at: {output_path}")

if __name__ == "__main__":
    generate_sample_excel()
