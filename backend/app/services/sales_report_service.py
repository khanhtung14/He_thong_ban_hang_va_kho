"""Sales report service for multi-dimensional sales analytics and Excel export (SCRUM-72)."""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, select
from sqlalchemy.orm import Session, joinedload, selectinload

from app.api.v1.endpoints.rbac import AuthenticatedUser, normalize_role
from app.models.models import (
    Category,
    Customer,
    Order,
    OrderItem,
    OrderStatus,
    Product,
    RoleCode,
    Territory,
    User,
)

VALID_ORDER_STATUSES = [
    OrderStatus.APPROVED,
    OrderStatus.PROCESSING,
    OrderStatus.COMPLETED,
    "APPROVED",
    "PROCESSING",
    "COMPLETED",
]


def is_sales_manager_or_admin(user: AuthenticatedUser) -> bool:
    """Check if the user is a Sales Manager or Admin permitted to view gross margin."""
    norm = normalize_role(user.role)
    return norm in (RoleCode.SALES_MANAGER.value, RoleCode.ADMIN.value, "Admin", "Sales Manager")


def parse_date_param(val: Any) -> Optional[date]:
    """Parse date string into date object."""
    if not isinstance(val, str):
        return None
    val = val.strip()
    if not val:
        return None
    # ISO string YYYY-MM-DD or YYYY-MM-DDTHH:MM:SS
    if "T" in val:
        val = val.split("T")[0]
    elif " " in val:
        val = val.split(" ")[0]
    try:
        return datetime.strptime(val, "%Y-%m-%d").date()
    except ValueError:
        try:
            return datetime.strptime(val, "%d/%m/%Y").date()
        except ValueError:
            return None


def resolve_period_dates(
    from_date_str: Optional[str] = None,
    to_date_str: Optional[str] = None,
    period: Optional[str] = None,
) -> Tuple[Optional[datetime], Optional[datetime], Optional[date], Optional[date]]:
    """Resolve start and end datetimes based on query params."""
    parsed_from = parse_date_param(from_date_str)
    parsed_to = parse_date_param(to_date_str)

    today = date.today()

    if period:
        period_clean = period.strip().lower()
        if period_clean == "today":
            parsed_from = today
            parsed_to = today
        elif period_clean == "this_month":
            parsed_from = date(today.year, today.month, 1)
            # Last day of current month
            next_month = date(today.year + (1 if today.month == 12 else 0), 1 if today.month == 12 else today.month + 1, 1)
            parsed_to = next_month - timedelta(days=1)
        elif period_clean == "last_month":
            last_month_end = date(today.year, today.month, 1) - timedelta(days=1)
            parsed_from = date(last_month_end.year, last_month_end.month, 1)
            parsed_to = last_month_end
        elif period_clean == "this_quarter":
            quarter = (today.month - 1) // 3
            parsed_from = date(today.year, quarter * 3 + 1, 1)
            if quarter == 3:
                parsed_to = date(today.year, 12, 31)
            else:
                parsed_to = date(today.year, (quarter + 1) * 3 + 1, 1) - timedelta(days=1)
        elif period_clean == "this_year":
            parsed_from = date(today.year, 1, 1)
            parsed_to = date(today.year, 12, 31)

    start_dt: Optional[datetime] = None
    end_dt: Optional[datetime] = None

    if parsed_from:
        start_dt = datetime.combine(parsed_from, time.min)
    if parsed_to:
        end_dt = datetime.combine(parsed_to, time.max)

    return start_dt, end_dt, parsed_from, parsed_to


def resolve_comparison_period(
    start_dt: Optional[datetime],
    end_dt: Optional[datetime],
    parsed_from: Optional[date],
    parsed_to: Optional[date],
    compare_to: Optional[str] = None,
    compare_from_str: Optional[str] = None,
    compare_to_str: Optional[str] = None,
) -> Tuple[Optional[datetime], Optional[datetime], Optional[date], Optional[date]]:
    """Determine previous period comparison dates."""
    c_from = parse_date_param(compare_from_str)
    c_to = parse_date_param(compare_to_str)

    if c_from and c_to:
        return (
            datetime.combine(c_from, time.min),
            datetime.combine(c_to, time.max),
            c_from,
            c_to,
        )

    if not compare_to:
        return None, None, None, None

    cmp_mode = compare_to.strip().lower()

    if not parsed_from or not parsed_to:
        # Fallback if no main period given
        today = date.today()
        parsed_from = date(today.year, today.month, 1)
        parsed_to = today

    if cmp_mode in ("previous_year", "same_period_last_year", "last_year"):
        try:
            prev_from = parsed_from.replace(year=parsed_from.year - 1)
        except ValueError:
            prev_from = parsed_from - timedelta(days=365)
        try:
            prev_to = parsed_to.replace(year=parsed_to.year - 1)
        except ValueError:
            prev_to = parsed_to - timedelta(days=365)
    else:
        # Default: previous_period (period of same length immediately preceding)
        num_days = (parsed_to - parsed_from).days + 1
        # Check if full month
        if parsed_from.day == 1 and (parsed_to + timedelta(days=1)).day == 1:
            prev_to = parsed_from - timedelta(days=1)
            prev_from = date(prev_to.year, prev_to.month, 1)
        else:
            prev_to = parsed_from - timedelta(days=1)
            prev_from = prev_to - timedelta(days=num_days - 1)

    return (
        datetime.combine(prev_from, time.min),
        datetime.combine(prev_to, time.max),
        prev_from,
        prev_to,
    )


def normalize_group_by(dim: Any) -> str:
    """Normalize grouping dimension to one of the 5 supported dimensions."""
    if not isinstance(dim, str):
        return "product"
    d = dim.strip().lower().replace("-", "_")
    if d in ("staff", "sales_rep", "salesperson", "rep", "employee", "nhan_vien"):
        return "staff"
    if d in ("region", "territory", "area", "zone", "khu_vuc", "dia_ban"):
        return "region"
    if d in ("customer", "dealer", "agency", "client", "dai_ly", "khach_hang"):
        return "customer"
    if d in ("category", "product_category", "cat", "nhom_hang"):
        return "category"
    if d in ("sku", "product", "item", "mat_hang", "san_pham"):
        return "product"
    return d


def query_and_aggregate_sales(
    db: Session,
    group_by: str,
    start_dt: Optional[datetime],
    end_dt: Optional[datetime],
) -> Dict[Any, Dict[str, Any]]:
    """Query orders within timeframe and aggregate revenue, volume, and COGS by dimension."""
    query = (
        db.query(Order)
        .options(
            joinedload(Order.customer).joinedload(Customer.sales_rep),
            selectinload(Order.items),
        )
        .filter(Order.status.in_(VALID_ORDER_STATUSES))
    )

    if start_dt is not None:
        query = query.filter(Order.created_at >= start_dt)
    if end_dt is not None:
        query = query.filter(Order.created_at <= end_dt)

    orders = query.all()

    # Pre-fetch lookup tables for performance and safety
    all_products = {p.id: p for p in db.query(Product).all()}
    products_by_sku = {p.sku.upper(): p for p in all_products.values() if p.sku}
    all_categories = {c.id: c for c in db.query(Category).all()}
    all_users = {u.id: u for u in db.query(User).all()}
    all_territories = {t.id: t for t in db.query(Territory).all()}

    # aggregation map: group_key -> {"id": ..., "code": ..., "name": ..., "revenue": ..., "quantity": ..., "cogs": ...}
    aggregated: Dict[Any, Dict[str, Any]] = {}

    for order in orders:
        cust = order.customer
        # Staff assignment
        sales_rep_user: Optional[User] = None
        if order.created_by_id and order.created_by_id in all_users:
            sales_rep_user = all_users[order.created_by_id]
        elif cust and cust.sales_rep_id and cust.sales_rep_id in all_users:
            sales_rep_user = all_users[cust.sales_rep_id]

        # Region assignment
        territory = all_territories.get(cust.territory_id) if (cust and cust.territory_id) else None
        region_name = "Chưa phân vùng"
        region_code = "OTHER"
        region_id = 0
        if cust:
            if cust.region and cust.region.strip():
                region_name = cust.region.strip()
                region_code = territory.code if territory else region_name.upper().replace(" ", "_")
                region_id = territory.id if territory else 0
            elif territory:
                region_name = territory.name
                region_code = territory.code or region_name
                region_id = territory.id

        # Customer assignment
        customer_id = cust.id if cust else 0
        customer_code = cust.code if cust else "UNKNOWN"
        customer_name = cust.name if cust else "Khách vãng lai"

        for line in order.items:
            product = all_products.get(line.product_id) if line.product_id else None
            if not product and line.sku:
                product = products_by_sku.get(line.sku.upper())

            sku = line.sku or (product.sku if product else "UNKNOWN")
            product_name = line.product_name or (product.name if product else sku)
            product_id = product.id if product else (line.product_id or 0)

            # Category
            category = all_categories.get(product.category_id) if (product and product.category_id) else None
            category_id = category.id if category else 0
            category_code = category.code if category else "OTHER"
            category_name = category.name if category else "Chưa phân loại"

            # Line calculations
            qty = line.quantity or 0
            line_rev = 0.0
            if line.line_total is not None and line.line_total > 0:
                line_rev = float(line.line_total)
            elif line.line_subtotal is not None and line.line_subtotal > 0:
                line_rev = float(line.line_subtotal) - float(line.discount_amount or 0.0)
            else:
                line_rev = (float(line.quantity or 0) * float(line.unit_price or 0.0)) - float(line.discount_amount or 0.0)

            # COGS
            cost_unit = float(product.cost_price or 0.0) if product else 0.0
            line_cogs = qty * cost_unit

            # Determine key based on dimension
            if group_by == "staff":
                gid = sales_rep_user.id if sales_rep_user else 0
                gcode = sales_rep_user.username if sales_rep_user else "UNASSIGNED"
                gname = sales_rep_user.full_name or sales_rep_user.username if sales_rep_user else "Chưa phân công"
                gkey = f"staff_{gid}"
            elif group_by == "region":
                gid = region_id
                gcode = region_code
                gname = region_name
                gkey = f"region_{gname}"
            elif group_by == "customer":
                gid = customer_id
                gcode = customer_code
                gname = customer_name
                gkey = f"customer_{gid}"
            elif group_by == "category":
                gid = category_id
                gcode = category_code
                gname = category_name
                gkey = f"category_{gid}"
            else:  # product / sku
                gid = product_id
                gcode = sku
                gname = product_name
                gkey = f"product_{sku}"

            if gkey not in aggregated:
                aggregated[gkey] = {
                    "id": gid,
                    "code": gcode,
                    "name": gname,
                    "revenue": 0.0,
                    "quantity": 0,
                    "cogs": 0.0,
                }

            aggregated[gkey]["revenue"] += line_rev
            aggregated[gkey]["quantity"] += qty
            aggregated[gkey]["cogs"] += line_cogs

    return aggregated


def calculate_growth_rate(current: float, previous: float) -> float:
    """Calculate percentage growth rate between two values."""
    if previous > 0:
        return round(((current - previous) / previous) * 100.0, 2)
    elif current > 0:
        return 100.0
    return 0.0


def build_sales_analytics_report(
    db: Session,
    user: AuthenticatedUser,
    group_by_param: Optional[str] = None,
    from_date_str: Optional[str] = None,
    to_date_str: Optional[str] = None,
    period: Optional[str] = None,
    compare_to: Optional[str] = None,
    compare_from_str: Optional[str] = None,
    compare_to_str: Optional[str] = None,
) -> Dict[str, Any]:
    """Build multi-dimensional sales report with RBAC and period comparison."""
    group_by = normalize_group_by(group_by_param)
    start_dt, end_dt, p_from, p_to = resolve_period_dates(from_date_str, to_date_str, period)
    cmp_start_dt, cmp_end_dt, cmp_from, cmp_to = resolve_comparison_period(
        start_dt, end_dt, p_from, p_to, compare_to, compare_from_str, compare_to_str
    )

    can_view_margins = is_sales_manager_or_admin(user)

    # 1. Main Period Aggregation
    current_agg = query_and_aggregate_sales(db, group_by, start_dt, end_dt)

    # 2. Comparison Period Aggregation (if comparison enabled)
    has_comparison = cmp_start_dt is not None and cmp_end_dt is not None
    cmp_agg: Dict[Any, Dict[str, Any]] = {}
    if has_comparison:
        cmp_agg = query_and_aggregate_sales(db, group_by, cmp_start_dt, cmp_end_dt)

    # Combine all keys from both current and comparison
    all_keys = set(current_agg.keys()) | (set(cmp_agg.keys()) if has_comparison else set())

    items: List[Dict[str, Any]] = []
    total_rev = 0.0
    total_qty = 0
    total_cogs = 0.0
    total_prev_rev = 0.0
    total_prev_qty = 0

    for key in sorted(all_keys):
        c_item = current_agg.get(key, {"id": 0, "code": "", "name": "", "revenue": 0.0, "quantity": 0, "cogs": 0.0})
        p_item = cmp_agg.get(key, {"revenue": 0.0, "quantity": 0}) if has_comparison else {"revenue": 0.0, "quantity": 0}

        gid = c_item["id"] if c_item["code"] else p_item.get("id", 0)
        gcode = c_item["code"] or p_item.get("code", "")
        gname = c_item["name"] or p_item.get("name", "")
        rev = round(c_item["revenue"], 2)
        qty = int(c_item["quantity"])
        cogs = round(c_item["cogs"], 2)
        gross_profit = round(rev - cogs, 2)
        margin = round(((rev - cogs) / rev * 100.0), 2) if rev > 0 else 0.0

        prev_rev = round(p_item["revenue"], 2) if has_comparison else 0.0
        prev_qty = int(p_item["quantity"]) if has_comparison else 0
        rev_growth = calculate_growth_rate(rev, prev_rev) if has_comparison else 0.0
        qty_growth = calculate_growth_rate(qty, prev_qty) if has_comparison else 0.0

        total_rev += rev
        total_qty += qty
        total_cogs += cogs
        total_prev_rev += prev_rev
        total_prev_qty += prev_qty

        item_dict: Dict[str, Any] = {
            "id": gid,
            "code": gcode,
            "name": gname,
            "revenue": rev,
            "quantity": qty,
            # Comparison metrics
            "previous_revenue": prev_rev if has_comparison else None,
            "previous_quantity": prev_qty if has_comparison else None,
            "revenue_growth_rate": rev_growth if has_comparison else None,
            "quantity_growth_rate": qty_growth if has_comparison else None,
        }

        # RBAC: Only Manager or Admin can view COGS & Gross Margin
        if can_view_margins:
            item_dict["cogs"] = cogs
            item_dict["gross_profit"] = gross_profit
            item_dict["gross_margin"] = margin
            item_dict["gross_margin_percentage"] = f"{margin:.1f}%"
        else:
            item_dict["cogs"] = None
            item_dict["gross_profit"] = None
            item_dict["gross_margin"] = None
            item_dict["gross_margin_percentage"] = None

        items.append(item_dict)

    # Sort by revenue descending
    items.sort(key=lambda x: x["revenue"], reverse=True)

    # Total calculations
    total_rev = round(total_rev, 2)
    total_cogs = round(total_cogs, 2)
    total_gross_profit = round(total_rev - total_cogs, 2)
    total_gross_margin = round(((total_rev - total_cogs) / total_rev * 100.0), 2) if total_rev > 0 else 0.0
    total_rev_growth = calculate_growth_rate(total_rev, total_prev_rev) if has_comparison else 0.0
    total_qty_growth = calculate_growth_rate(total_qty, total_prev_qty) if has_comparison else 0.0

    summary: Dict[str, Any] = {
        "total_revenue": total_rev,
        "total_quantity": total_qty,
        "previous_revenue": round(total_prev_rev, 2) if has_comparison else None,
        "previous_quantity": total_prev_qty if has_comparison else None,
        "revenue_growth_rate": total_rev_growth if has_comparison else None,
        "quantity_growth_rate": total_qty_growth if has_comparison else None,
    }

    if can_view_margins:
        summary["total_cogs"] = total_cogs
        summary["gross_profit"] = total_gross_profit
        summary["gross_margin"] = total_gross_margin
        summary["gross_margin_percentage"] = f"{total_gross_margin:.1f}%"
    else:
        summary["total_cogs"] = None
        summary["gross_profit"] = None
        summary["gross_margin"] = None
        summary["gross_margin_percentage"] = None

    period_info = {
        "from_date": p_from.strftime("%Y-%m-%d") if p_from else None,
        "to_date": p_to.strftime("%Y-%m-%d") if p_to else None,
    }

    comparison_info = None
    if has_comparison:
        comparison_info = {
            "compare_to": compare_to or "custom",
            "from_date": cmp_from.strftime("%Y-%m-%d") if cmp_from else None,
            "to_date": cmp_to.strftime("%Y-%m-%d") if cmp_to else None,
        }

    return {
        "success": True,
        "code": 200,
        "message": "Lấy báo cáo doanh số đa chiều thành công",
        "data": {
            "group_by": group_by,
            "period": period_info,
            "comparison": comparison_info,
            "summary": summary,
            "items": items,
            "details": items,  # Alias for backward compatibility
        },
    }


def generate_sales_analytics_excel(
    report_data: Dict[str, Any],
    user: AuthenticatedUser,
) -> bytes:
    """Export the multi-dimensional sales analytics report to a beautifully formatted Excel sheet."""
    data = report_data.get("data", {})
    group_by = data.get("group_by", "category")
    summary = data.get("summary", {})
    items = data.get("items", [])
    period = data.get("period", {})
    comparison = data.get("comparison")
    can_view_margins = is_sales_manager_or_admin(user)
    has_comparison = comparison is not None

    wb = Workbook()
    ws = wb.active
    ws.title = "Doanh Số Đa Chiều"
    ws.views.sheetView[0].showGridLines = True

    # Styling Token constants
    NAVY = "1F497D"
    WHITE = "FFFFFF"
    GRAY_HEADER = "F2F2F2"
    TOTAL_ROW_FILL = "D9E1F2"

    font_title = Font(name="Calibri", size=15, bold=True, color="1F497D")
    font_sub = Font(name="Calibri", size=10, italic=True, color="595959")
    font_header = Font(name="Calibri", size=11, bold=True, color=WHITE)
    font_total = Font(name="Calibri", size=11, bold=True, color="000000")
    font_data = Font(name="Calibri", size=11)

    fill_header = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    fill_total = PatternFill(start_color=TOTAL_ROW_FILL, end_color=TOTAL_ROW_FILL, fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )
    total_border = Border(
        left=Side(style="thin", color="1F497D"),
        right=Side(style="thin", color="1F497D"),
        top=Side(style="thin", color="1F497D"),
        bottom=Side(style="double", color="1F497D"),
    )

    # 1. Header Information Block
    dim_titles = {
        "staff": "Nhân viên kinh doanh",
        "region": "Khu vực / Địa bàn",
        "customer": "Đại lý / Khách hàng",
        "category": "Nhóm hàng",
        "product": "Mặt hàng cụ thể (SKU)",
    }
    dim_label = dim_titles.get(group_by, group_by.capitalize())

    ws["A1"] = "HỆ THỐNG QUẢN LÝ BÁN HÀNG & KHO - BÁO CÁO DOANH SỐ ĐA CHIỀU"
    ws["A1"].font = font_title

    p_str = f"{period.get('from_date') or 'Tất cả'} đến {period.get('to_date') or 'Hiện tại'}"
    ws["A2"] = f"Tiêu chí gom nhóm: {dim_label} | Kỳ báo cáo: {p_str}"
    ws["A2"].font = font_sub

    cmp_str = f"So sánh với: {comparison.get('from_date')} đến {comparison.get('to_date')}" if has_comparison else "So sánh: Không"
    user_str = f"Người xuất: {user.full_name or user.username} ({user.role}) | Ngày tạo: {date.today().strftime('%d/%m/%Y')}"
    ws["A3"] = f"{user_str} | {cmp_str}"
    ws["A3"].font = font_sub

    # 2. Table Column Definitions
    headers = ["STT", "Mã", f"Tên {dim_label}", "Doanh số (VNĐ)", "Sản lượng bán"]

    if can_view_margins:
        headers.extend(["Giá vốn COGS (VNĐ)", "Lợi nhuận gộp (VNĐ)", "Biên lợi nhuận (%)"])

    if has_comparison:
        headers.extend(["Doanh số kỳ trước (VNĐ)", "Sản lượng kỳ trước", "Tăng trưởng DS (%)", "Tăng trưởng SL (%)"])

    # Write Table Header at Row 5
    header_row = 5
    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[header_row].height = 28

    # 3. Write Data Rows
    current_row = header_row + 1
    for idx, item in enumerate(items, 1):
        row_vals = [
            idx,
            item["code"],
            item["name"],
            item["revenue"],
            item["quantity"],
        ]

        if can_view_margins:
            row_vals.extend([
                item["cogs"] if item["cogs"] is not None else 0.0,
                item["gross_profit"] if item["gross_profit"] is not None else 0.0,
                (item["gross_margin"] or 0.0) / 100.0,
            ])

        if has_comparison:
            row_vals.extend([
                item["previous_revenue"] if item["previous_revenue"] is not None else 0.0,
                item["previous_quantity"] if item["previous_quantity"] is not None else 0,
                (item["revenue_growth_rate"] or 0.0) / 100.0,
                (item["quantity_growth_rate"] or 0.0) / 100.0,
            ])

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = thin_border

            # Number formatting
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in (2, 3):
                cell.alignment = Alignment(horizontal="left")
            else:
                cell.alignment = Alignment(horizontal="right")

            # Check format types
            header_name = headers[col_idx - 1]
            if "(VNĐ)" in header_name:
                cell.number_format = "#,##0"
            elif "(%)" in header_name:
                cell.number_format = "0.0%"
            elif "Sản lượng" in header_name:
                cell.number_format = "#,##0"

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # 4. Total Row
    total_row_idx = current_row
    ws.cell(row=total_row_idx, column=1, value="")
    ws.cell(row=total_row_idx, column=2, value="TỔNG CỘNG")
    ws.cell(row=total_row_idx, column=3, value="")

    total_vals: Dict[str, Any] = {
        "Doanh số (VNĐ)": summary.get("total_revenue", 0.0),
        "Sản lượng bán": summary.get("total_quantity", 0),
    }

    if can_view_margins:
        total_vals["Giá vốn COGS (VNĐ)"] = summary.get("total_cogs", 0.0)
        total_vals["Lợi nhuận gộp (VNĐ)"] = summary.get("gross_profit", 0.0)
        total_vals["Biên lợi nhuận (%)"] = (summary.get("gross_margin") or 0.0) / 100.0

    if has_comparison:
        total_vals["Doanh số kỳ trước (VNĐ)"] = summary.get("previous_revenue", 0.0) or 0.0
        total_vals["Sản lượng kỳ trước"] = summary.get("previous_quantity", 0) or 0
        total_vals["Tăng trưởng DS (%)"] = (summary.get("revenue_growth_rate") or 0.0) / 100.0
        total_vals["Tăng trưởng SL (%)"] = (summary.get("quantity_growth_rate") or 0.0) / 100.0

    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=total_row_idx, column=col_idx)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = total_border

        if h_text in total_vals:
            cell.value = total_vals[h_text]
            cell.alignment = Alignment(horizontal="right")
            if "(VNĐ)" in h_text:
                cell.number_format = "#,##0"
            elif "(%)" in h_text:
                cell.number_format = "0.0%"
            elif "Sản lượng" in h_text:
                cell.number_format = "#,##0"
        elif col_idx in (1, 2, 3):
            cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "left")

    ws.row_dimensions[total_row_idx].height = 24

    # 5. Auto adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip title rows in calculating column width
            if cell.row < 5:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Specific tweaks for common columns
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
