from io import BytesIO
import logging
import re
import secrets
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models.models import (
    AccountAuditLog,
    AccountStatus,
    Role,
    Territory,
    User,
    Warehouse,
    seed_default_roles,
)
from app.services.email_service import email_service
from app.services.users_service import (
    generate_temporary_password,
    hash_password,
)

logger = logging.getLogger("import_users")
router = APIRouter()

DEFAULT_PASSWORD = "User@123456"

REQUIRED_HEADERS = [
    "username",
    "full_name",
    "email",
    "role",
]

OPTIONAL_HEADERS = [
    "phone",
    "warehouse",
    "region",
]

ALL_HEADERS = REQUIRED_HEADERS + OPTIONAL_HEADERS

# Bảng mapping vai trò hệ thống: hỗ trợ cả mã code lẫn tên hiển thị tiếng Việt
ROLE_MAPPING = {
    # sales / Nhân viên kinh doanh
    "sales": "sales",
    "nhân viên kinh doanh": "sales",
    "nhan vien kinh doanh": "sales",
    "sales_rep": "sales",

    # sales_manager / Quản lý kinh doanh
    "sales_manager": "sales_manager",
    "quản lý kinh doanh": "sales_manager",
    "quan ly kinh doanh": "sales_manager",

    # warehouse_staff / Nhân viên kho
    "warehouse_staff": "warehouse_staff",
    "nhân viên kho": "warehouse_staff",
    "nhan vien kho": "warehouse_staff",
    "warehouse": "warehouse_staff",

    # warehouse_manager / Quản lý kho
    "warehouse_manager": "warehouse_manager",
    "quản lý kho": "warehouse_manager",
    "quan ly kho": "warehouse_manager",
    "wh_manager": "warehouse_manager",

    # accountant / Kế toán
    "accountant": "accountant",
    "kế toán": "accountant",
    "ke toan": "accountant",

    # admin / Quản trị hệ thống
    "admin": "admin",
    "quản trị hệ thống": "admin",
    "quan tri he thong": "admin",

    # agency / Đại lý
    "agency": "agency",
    "đại lý": "agency",
    "dai ly": "agency",
    "customer": "agency",
}

# Mapping từ chuẩn vai trò sang Role.code trong cơ sở dữ liệu
ROLE_TO_DB_CODE = {
    "sales": "SALES_REP",
    "sales_manager": "SALES_MANAGER",
    "warehouse_staff": "WAREHOUSE",
    "warehouse_manager": "WH_MANAGER",
    "accountant": "ACCOUNTANT",
    "admin": "ADMIN",
    "agency": "CUSTOMER",
}


def create_excel_template():
    """Tạo file Excel mẫu để quản trị viên tải về."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sach nguoi dung"

    sheet.append([
        "username",
        "full_name",
        "email",
        "role",
        "phone",
        "warehouse",
        "region",
    ])
    sheet.append([
        "nguyenvana",
        "Nguyen Van A",
        "vana@example.com",
        "sales",
        "0912345678",
        "Kho Tổng Hà Nội",
        "Miền Bắc",
    ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def read_excel_file(file_content):
    """Đọc dữ liệu từ file Excel và kiểm tra tiêu đề cột."""
    try:
        workbook = load_workbook(
            BytesIO(file_content),
            read_only=True,
            data_only=True,
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Không đọc được file Excel. Vui lòng kiểm tra lại file.",
        )

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    headers = next(rows, None)

    if not headers:
        raise HTTPException(
            status_code=400,
            detail="File Excel đang trống.",
        )

    headers = [
        str(value).strip().lower() if value is not None else ""
        for value in headers
    ]

    missing_headers = [
        header
        for header in REQUIRED_HEADERS
        if header not in headers
    ]

    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "File Excel thiếu cột bắt buộc.",
                "missing_headers": missing_headers,
            },
        )

    header_positions = {
        header: headers.index(header)
        for header in ALL_HEADERS
        if header in headers
    }

    result = []

    for row_number, row in enumerate(rows, start=2):
        user_data = {}

        for header in ALL_HEADERS:
            position = header_positions.get(header)
            value = (
                row[position]
                if (position is not None and position < len(row))
                else None
            )

            user_data[header] = (
                str(value).strip() if value is not None else ""
            )

        # Bỏ qua dòng hoàn toàn trống.
        if not any(user_data.values()):
            continue

        result.append({
            "row_number": row_number,
            "data": user_data,
        })

    workbook.close()
    return result


def validate_rows(rows, db: Optional[Session] = None):
    """Kiểm tra dữ liệu từng dòng, đối chiếu với cơ sở dữ liệu thật."""
    results = []

    seen_usernames = set()
    seen_emails = set()
    seen_phones = set()

    existing_usernames = set()
    existing_emails = set()
    existing_phones = set()
    active_warehouses = []
    all_territories = []

    if db is not None:
        db_usernames = db.query(User.username).filter(User.username.isnot(None)).all()
        existing_usernames = {u[0].strip().lower() for u in db_usernames if u[0]}

        db_emails = db.query(User.email).filter(User.email.isnot(None)).all()
        existing_emails = {e[0].strip().lower() for e in db_emails if e[0]}

        db_phones = db.query(User.phone).filter(User.phone.isnot(None)).all()
        existing_phones = {p[0].strip() for p in db_phones if p[0]}

        active_warehouses = db.query(Warehouse).filter(Warehouse.is_active.is_(True)).all()
        all_territories = db.query(Territory).all()

    for item in rows:
        row_number = item["row_number"]
        data = item["data"]
        errors = []

        username = str(data.get("username", "") or "").strip()
        full_name = str(data.get("full_name", "") or "").strip()
        email = str(data.get("email", "") or "").strip()
        raw_role = str(data.get("role", "") or "").strip()
        phone = str(data.get("phone", "") or "").strip()
        warehouse = str(data.get("warehouse", "") or "").strip()
        region = str(data.get("region", "") or "").strip()

        # Chuẩn hóa dữ liệu đã trim
        data["username"] = username
        data["full_name"] = full_name
        data["email"] = email
        data["role"] = raw_role
        data["phone"] = phone
        data["warehouse"] = warehouse
        data["region"] = region

        if not username:
            errors.append("Thiếu tên đăng nhập.")

        if not full_name:
            errors.append("Thiếu họ và tên.")

        if not email:
            errors.append("Thiếu email.")
        elif not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            errors.append("Email không đúng định dạng.")

        if not raw_role:
            errors.append("Thiếu vai trò.")
        else:
            role_key = raw_role.lower()
            if role_key in ROLE_MAPPING:
                mapped_role = ROLE_MAPPING[role_key]
                data["role"] = mapped_role
            else:
                errors.append("Vai trò không hợp lệ trong hệ thống.")

        # Siết chặt ràng buộc kho theo vai trò (Conditional Validation):
        # Nếu role là warehouse_staff hoặc warehouse_manager: cột warehouse là bắt buộc
        if data.get("role") in ("warehouse_staff", "warehouse_manager"):
            if not warehouse:
                errors.append("Nhân sự bộ phận kho bắt buộc phải được gán Kho phụ trách.")

        # Kiểm tra sự tồn tại của kho phụ trách nếu có nhập
        if warehouse and db is not None:
            wh_match = any(
                warehouse.lower() == w.code.lower()
                or warehouse.lower() == w.name.lower()
                or warehouse.lower() in w.name.lower()
                or warehouse.lower() in w.code.lower()
                for w in active_warehouses
            )
            if not wh_match:
                errors.append(f"Kho phụ trách '{warehouse}' không tồn tại hoặc đã ngừng hoạt động trong hệ thống.")

        # Kiểm tra sự tồn tại của địa bàn hoạt động nếu có nhập
        if region and db is not None:
            terr_match = any(
                region.lower() == t.code.lower()
                or region.lower() == t.name.lower()
                or region.lower() in t.name.lower()
                or region.lower() in t.code.lower()
                for t in all_territories
            )
            if not terr_match:
                errors.append(f"Địa bàn hoạt động '{region}' không tồn tại trong hệ thống.")

        # Kiểm tra trùng tên đăng nhập
        if username:
            username_key = username.lower()
            if username_key in existing_usernames:
                errors.append("Tên đăng nhập đã tồn tại trong hệ thống.")
            elif username_key in seen_usernames:
                errors.append("Tên đăng nhập bị trùng trong file.")
            seen_usernames.add(username_key)

        # Kiểm tra trùng email
        if email:
            email_key = email.lower()
            if email_key in existing_emails:
                errors.append("Email đã tồn tại trong hệ thống.")
            elif email_key in seen_emails:
                errors.append("Email bị trùng trong file.")
            seen_emails.add(email_key)

        # Kiểm tra trùng số điện thoại (nếu có nhập)
        if phone:
            if phone in existing_phones:
                errors.append("Số điện thoại đã tồn tại trong hệ thống.")
            elif phone in seen_phones:
                errors.append("Số điện thoại bị trùng trong file.")
            seen_phones.add(phone)

        results.append({
            "row_number": row_number,
            "data": data,
            "valid": len(errors) == 0,
            "errors": errors,
        })

    return results


@router.get("/users/import/template")
@router.get("/api/v1/users/import/template")
def download_users_template():
    """Tải file Excel mẫu."""
    excel_file = create_excel_template()

    return StreamingResponse(
        excel_file,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                "attachment; filename=users_template.xlsx"
            )
        },
    )


@router.post("/users/import/preview")
@router.post("/api/v1/users/import/preview")
async def preview_users_import(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Xem trước dữ liệu và đối chiếu kiểm tra lỗi với cơ sở dữ liệu thật."""
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Vui lòng tải lên file Excel có đuôi .xlsx.",
        )

    content = await file.read()
    rows = read_excel_file(content)
    results = validate_rows(rows, db=db)

    valid_count = sum(1 for item in results if item["valid"])
    error_count = len(results) - valid_count

    return {
        "total_rows": len(results),
        "valid_rows": valid_count,
        "error_rows": error_count,
        "rows": results,
    }


@router.post("/users/import")
@router.post("/api/v1/users/import")
async def import_users(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Nhập các dòng hợp lệ vào cơ sở dữ liệu thật và bỏ qua các dòng lỗi."""
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Vui lòng tải lên file Excel có đuôi .xlsx.",
        )

    content = await file.read()
    rows = read_excel_file(content)
    results = validate_rows(rows, db=db)

    # Đảm bảo seed các role chuẩn trước khi import
    seed_default_roles(db)

    created_count = 0
    skipped_count = 0
    report = []

    for item in results:
        if not item["valid"]:
            skipped_count += 1
            report.append({
                "row_number": item["row_number"],
                "status": "skipped",
                "email_sent": False,
                "message": "; ".join(item["errors"]),
                "errors": item["errors"],
            })
            continue

        user_data = item["data"]
        role_key = user_data.get("role", "")
        target_role_code = ROLE_TO_DB_CODE.get(role_key.lower(), role_key.upper())

        role_obj = (
            db.query(Role)
            .filter(
                or_(
                    func.upper(Role.code) == target_role_code.upper(),
                    func.upper(Role.code) == role_key.upper(),
                    func.lower(Role.name) == role_key.lower(),
                )
            )
            .first()
        )
        assigned_roles = [role_obj] if role_obj else []

        # Gán kho nếu có
        assigned_warehouses = []
        warehouse_name = user_data.get("warehouse", "").strip()
        if warehouse_name:
            wh_obj = (
                db.query(Warehouse)
                .filter(
                    Warehouse.is_active.is_(True),
                    or_(
                        Warehouse.name.ilike(f"%{warehouse_name}%"),
                        Warehouse.code.ilike(f"%{warehouse_name}%"),
                    ),
                )
                .first()
            )
            if wh_obj:
                assigned_warehouses.append(wh_obj)

        # Gán địa bàn nếu có
        assigned_territories = []
        region_name = user_data.get("region", "").strip()
        if region_name:
            territory_obj = (
                db.query(Territory)
                .filter(
                    or_(
                        Territory.name.ilike(f"%{region_name}%"),
                        Territory.code.ilike(f"%{region_name}%"),
                    )
                )
                .first()
            )
            if territory_obj:
                assigned_territories.append(territory_obj)

        # Sinh mật khẩu tạm thời an toàn và mã hóa bcrypt
        temp_password = generate_temporary_password(12)
        password_hash = hash_password(temp_password)

        new_user = User(
            username=user_data["username"],
            full_name=user_data["full_name"],
            email=user_data["email"],
            phone=user_data.get("phone") or None,
            password_hash=password_hash,
            is_active=True,
            status=AccountStatus.PENDING_ACTIVATION,
            must_change_password=True,
            roles=assigned_roles,
            warehouses=assigned_warehouses,
            territories=assigned_territories,
        )

        try:
            db.add(new_user)
            db.commit()
            db.refresh(new_user)

            try:
                audit_log = AccountAuditLog(
                    user_id=new_user.id,
                    action="IMPORT_USER_EXCEL",
                    reason=f"Nhập tài khoản {new_user.username} từ file Excel (chờ kích hoạt)",
                )
                db.add(audit_log)
                db.commit()
            except Exception:
                db.rollback()

            # Gửi email kích hoạt tài khoản
            activation_token = secrets.token_urlsafe(32)
            email_sent = False
            try:
                email_sent = email_service.send_activation_email(
                    to_email=new_user.email,
                    full_name=new_user.full_name,
                    username=new_user.username,
                    temp_password=temp_password,
                    login_url="/login",
                    activation_token=activation_token,
                    max_retries=2,
                )
            except Exception as mail_exc:
                logger.warning(
                    f"Lỗi khi gửi email kích hoạt cho {new_user.email}: {mail_exc}"
                )
                email_sent = False

            created_count += 1
            status_msg = (
                "Chờ kích hoạt (Đã gửi email)"
                if email_sent
                else "Tạo thành công nhưng chưa gửi được mail kích hoạt"
            )
            report.append({
                "row_number": item["row_number"],
                "status": "created",
                "email_sent": email_sent,
                "message": status_msg,
                "errors": [],
            })
        except Exception as exc:
            db.rollback()
            skipped_count += 1
            report.append({
                "row_number": item["row_number"],
                "status": "skipped",
                "email_sent": False,
                "message": f"Lỗi lưu cơ sở dữ liệu: {str(exc)}",
                "errors": [f"Lỗi lưu cơ sở dữ liệu: {str(exc)}"],
            })

    return {
        "total_rows": len(results),
        "created_count": created_count,
        "skipped_count": skipped_count,
        "report": report,
    }