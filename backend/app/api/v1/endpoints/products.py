"""Products endpoints with RBAC enforcement."""
from __future__ import annotations

from datetime import datetime
from typing import Any, List, Optional, Union
import io
import math
import re
import uuid
from zipfile import BadZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from fastapi import APIRouter, Depends, File, HTTPException, Path as PathParam, Query, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel, Field, field_validator, model_validator

from app.core.database import get_db
from app.models.models import (
    Product,
    Category,
    ProductUnit,
    ProductPriceHistory,
    RoleCode,
    OrderItem,
    WarehouseReceiptItem,
    WarehouseTransferItem,
    WarehouseAuditItem,
    DispatchItem,
    ReturnOrderItem,
    InventoryAdjustmentItem,
)
from app.api.v1.endpoints.rbac import (
    PERM_PRODUCTS_VIEW,
    PERM_PRODUCTS_MANAGE,
    PERM_PRODUCTS_UNIT_MANAGE,
    PERM_PRODUCTS_IMPORT,
    AuthenticatedUser,
    require_permissions,
    require_roles,
    sanitize_financial_data,
    can_view_financials,
)
from app.services.price_history_service import (
    get_price_history,
    parse_date_boundary,
    record_price_history,
    serialize_price_history,
)

class ProductUnitResponse(BaseModel):
    id: int
    unitName: str
    conversionRate: float
    barcode: Optional[str] = None

class ProductUnitRequest(BaseModel):
    unit_name: str = Field(min_length=1, max_length=100)
    conversion_rate: float = Field(gt=0, allow_inf_nan=False)
    is_base_unit: bool = False

    @field_validator("unit_name")
    @classmethod
    def clean_unit_name(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("Tên đơn vị không được để trống.")
        return value

class ProductUnitInput(BaseModel):
    unitName: str
    conversionRate: float = 1.0
    barcode: Optional[str] = None
    isBaseUnit: bool = False

class ProductCreateRequest(BaseModel):
    sku: str = Field(min_length=1, max_length=50)
    name: str = Field(min_length=1, max_length=255)
    categoryId: Optional[int] = None
    baseUnit: str = Field(min_length=1, max_length=20)
    manageByLot: bool = False
    minStock: int = 0
    basePrice: Optional[float] = 0.0
    costPrice: Optional[float] = 0.0
    imageUrl: Optional[str] = None
    status: str = "ACTIVE"
    units: Optional[list[ProductUnitInput]] = None

class ProductUpdateRequest(BaseModel):
    name: Optional[str] = None
    categoryId: Optional[int] = None
    baseUnit: Optional[str] = None
    manageByLot: Optional[bool] = None
    minStock: Optional[int] = None
    basePrice: Optional[float] = None
    costPrice: Optional[float] = None
    imageUrl: Optional[str] = None
    status: Optional[str] = None
    units: Optional[list[ProductUnitInput]] = None

class ProductStatusUpdateRequest(BaseModel):
    status: str = Field(pattern="^(ACTIVE|INACTIVE)$")

class ProductBase(BaseModel):
    id: int
    sku: str
    name: str
    categoryId: Optional[int] = None
    baseUnit: str
    manageByLot: bool
    minStock: int
    basePrice: Optional[float] = None
    costPrice: Optional[float] = None
    status: str
    imageUrl: Optional[str] = None
    units: Optional[list[ProductUnitResponse]] = None

class PaginationInfo(BaseModel):
    page: int
    limit: int
    totalRecords: int
    totalPages: int

class ProductListResponse(BaseModel):
    success: bool
    code: int
    message: str
    data: list[ProductBase]
    pagination: PaginationInfo

class ProductDetailResponse(BaseModel):
    success: bool
    code: int
    message: str
    data: ProductBase

router = APIRouter(prefix="/api/v1/products", tags=["Products"])

DEFAULT_MOCK_PRODUCTS: dict[str, dict[str, Any]] = {
    "SKU-001": {
        "sku": "SKU-001",
        "name": "Nước tăng lực Red Bull 250ml",
        "category": "Nước tăng lực có gas",
        "category_id": 3,
        "sale_price": 500000,
        "cost_price": 350000,
        "margin": "30.0%",
        "stock_available": 120,
        "unit": "Thùng 24 lon",
    },
    "SKU-002": {
        "sku": "SKU-002",
        "name": "Cà phê lon Highlands 235ml",
        "category": "Cà phê",
        "category_id": None,
        "sale_price": 240000,
        "cost_price": 180000,
        "margin": "25.0%",
        "stock_available": 85,
        "unit": "Thùng 24 lon",
    },
}

MOCK_PRODUCTS: dict[str, dict[str, Any]] = {k: v.copy() for k, v in DEFAULT_MOCK_PRODUCTS.items()}


def reset_mock_products() -> None:
    global MOCK_PRODUCTS
    MOCK_PRODUCTS.clear()
    for k, v in DEFAULT_MOCK_PRODUCTS.items():
        MOCK_PRODUCTS[k] = v.copy()


MAX_IMPORT_BYTES = 10 * 1024 * 1024  # 10 MB
IMPORT_COLUMNS = (
    "sku",
    "name",
    "category",
    "sale_price",
    "cost_price",
    "stock_available",
    "unit",
)
REQUIRED_IMPORT_COLUMNS = {"sku", "name", "category"}

HEADER_ALIASES: dict[str, str] = {
    "sku": "sku",
    "mã sku": "sku",
    "ma sku": "sku",
    "mã sản phẩm": "sku",
    "ma san pham": "sku",
    "mã": "sku",
    "ma": "sku",
    "name": "name",
    "tên sản phẩm": "name",
    "ten san pham": "name",
    "tên": "name",
    "ten": "name",
    "category": "category",
    "danh mục": "category",
    "danh muc": "category",
    "ngành hàng": "category",
    "nganh hang": "category",
    "loại sản phẩm": "category",
    "sale_price": "sale_price",
    "giá bán": "sale_price",
    "gia ban": "sale_price",
    "base_price": "sale_price",
    "cost_price": "cost_price",
    "giá vốn": "cost_price",
    "gia von": "cost_price",
    "stock_available": "stock_available",
    "tồn kho": "stock_available",
    "ton kho": "stock_available",
    "số lượng tồn": "stock_available",
    "số lượng": "stock_available",
    "so luong": "stock_available",
    "unit": "unit",
    "đơn vị tính": "unit",
    "don vi tinh": "unit",
    "đvt": "unit",
    "dvt": "unit",
    "đơn vị": "unit",
}

IMPORT_SESSIONS: dict[str, list[dict[str, Any]]] = {}


async def _read_import_file(file: UploadFile) -> bytes:
    if not file.filename or not (
        file.filename.lower().endswith(".xlsx") or file.filename.lower().endswith(".xls")
    ):
        raise HTTPException(status_code=415, detail="Chỉ hỗ trợ tệp Excel .xlsx hoặc .xls.")
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Tệp Excel vượt quá giới hạn 10 MB.")
    return content


def _row_to_import_item(row_number: int, values: dict[str, Any]) -> dict[str, Any]:
    errors: list[str] = []
    sku = str(values.get("sku") or "").strip().upper()
    name = str(values.get("name") or "").strip()
    category = str(values.get("category") or "").strip()

    if not sku:
        errors.append("SKU không được để trống.")
    if not name:
        errors.append("Tên sản phẩm không được để trống.")
    if not category:
        errors.append("Danh mục không được để trống.")

    item: dict[str, Any] = {"sku": sku, "name": name, "category": category}

    for col in ("sale_price", "cost_price", "stock_available"):
        raw_val = values.get(col)
        if raw_val is None or str(raw_val).strip() == "":
            continue
        try:
            num = float(str(raw_val).strip())
            if not math.isfinite(num) or num < 0 or (col == "stock_available" and not num.is_integer()):
                raise ValueError
            item[col] = int(num) if num.is_integer() else num
        except (TypeError, ValueError):
            suffix = " nguyên." if col == "stock_available" else "."
            errors.append(f"{col} phải là số không âm{suffix}")

    unit = str(values.get("unit") or "").strip()
    item["unit"] = unit or "Cái"

    return {
        "row": row_number,
        "row_number": row_number,
        "sku": sku,
        "name": name,
        "category": category,
        "unit": item["unit"],
        "sale_price": item.get("sale_price"),
        "cost_price": item.get("cost_price"),
        "stock_available": item.get("stock_available"),
        "item": item,
        "errors": errors,
    }


def _parse_import_workbook(content: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="Tệp không phải Excel .xlsx hợp lệ.") from exc

    try:
        worksheet = workbook.active
        rows_iter = worksheet.iter_rows(values_only=True)
        headers = next(rows_iter, None)
        if not headers:
            raise HTTPException(status_code=422, detail="Tệp Excel không có dòng tiêu đề.")

        normalized_headers: list[str] = []
        for val in headers:
            raw = str(val or "").strip().lower()
            mapped = HEADER_ALIASES.get(raw, raw)
            normalized_headers.append(mapped)

        nonempty_headers = [h for h in normalized_headers if h]
        if len(nonempty_headers) != len(set(nonempty_headers)):
            raise HTTPException(status_code=422, detail="Tệp Excel có tên cột bị trùng.")

        missing_columns = REQUIRED_IMPORT_COLUMNS - set(nonempty_headers)
        if missing_columns:
            raise HTTPException(
                status_code=422,
                detail=f"Thiếu cột bắt buộc: {', '.join(sorted(missing_columns))}.",
            )

        parsed_rows: list[dict[str, Any]] = []
        for row_number, cells in enumerate(rows_iter, start=2):
            if not any(value is not None and str(value).strip() for value in cells):
                continue
            values = {
                header: cells[index] if index < len(cells) else None
                for index, header in enumerate(normalized_headers)
                if header in IMPORT_COLUMNS
            }
            parsed_rows.append(_row_to_import_item(row_number, values))

        if not parsed_rows:
            raise HTTPException(status_code=422, detail="Tệp Excel không có dữ liệu sản phẩm.")
        return parsed_rows
    finally:
        workbook.close()


def _get_existing_skus(db: Session) -> set[str]:
    skus = {s[0].upper() for s in db.query(Product.sku).all() if s[0]}
    skus |= {k.upper() for k in MOCK_PRODUCTS.keys()}
    return skus


def _annotate_import_rows(rows: list[dict[str, Any]], existing_skus: set[str]) -> list[dict[str, Any]]:
    seen_skus: set[str] = set()
    result: list[dict[str, Any]] = []

    for r in rows:
        row = dict(r)
        errors = list(row.get("errors", []))
        sku = row.get("sku", "").upper()

        if sku and sku in seen_skus:
            errors.append("SKU bị lặp trong tệp Excel.")
        if sku:
            seen_skus.add(sku)

        if errors:
            action = "error"
            action_type = "ERROR"
            msg = "; ".join(errors)
            err_msg = msg
        elif sku in existing_skus:
            action = "update"
            action_type = "UPDATE"
            msg = "SKU đã tồn tại trong hệ thống, dữ liệu sẽ được cập nhật."
            err_msg = None
        else:
            action = "create"
            action_type = "CREATE"
            msg = "Sản phẩm sẽ được tạo mới."
            err_msg = None

        row["action"] = action
        row["action_type"] = action_type
        row["message"] = msg
        row["error_message"] = err_msg
        row.pop("errors", None)
        result.append(row)

    return result


def _apply_import_rows(
    rows: list[dict[str, Any]],
    db: Session,
    changed_by: Optional[str] = None,
) -> dict[str, Any]:
    created_count = 0
    updated_count = 0
    category_cache: dict[str, Category] = {}

    for row in rows:
        item = row.get("item") or row
        sku = str(item.get("sku") or "").strip().upper()
        if not sku:
            continue

        category_name = str(item.get("category") or "").strip()
        category_obj: Optional[Category] = None
        if category_name:
            if category_name in category_cache:
                category_obj = category_cache[category_name]
            else:
                category_obj = db.query(Category).filter(
                    (Category.name.ilike(category_name)) | (Category.code.ilike(category_name))
                ).first()
                if not category_obj:
                    clean_code = re.sub(r"[^A-Za-z0-9]", "", category_name).upper()[:15] or "CAT"
                    existing_code = db.query(Category).filter(Category.code == clean_code).first()
                    if existing_code:
                        clean_code = f"{clean_code[:10]}_{uuid.uuid4().hex[:4].upper()}"
                    category_obj = Category(code=clean_code, name=category_name)
                    db.add(category_obj)
                    db.flush()
                category_cache[category_name] = category_obj

        existing_product = db.query(Product).filter(Product.sku.ilike(sku)).first()
        if existing_product:
            updated_count += 1
            if item.get("name"):
                existing_product.name = str(item["name"]).strip()
            if category_obj:
                existing_product.category_id = category_obj.id
            if item.get("sale_price") is not None:
                new_p = float(item["sale_price"])
                old_p = existing_product.base_price
                if old_p != new_p:
                    existing_product._suppress_auto_price_history = True
                    existing_product.base_price = new_p
                    record_price_history(
                        db=db,
                        product_id=existing_product.id,
                        old_price=old_p,
                        new_price=new_p,
                        changed_by=changed_by or "System",
                        note="Cập nhật giá bán qua import Excel",
                        reason="Cập nhật giá bán qua import Excel",
                    )
            if item.get("cost_price") is not None:
                existing_product.cost_price = float(item["cost_price"])
            if item.get("unit"):
                existing_product.base_unit = str(item["unit"]).strip()
        elif sku in MOCK_PRODUCTS:
            updated_count += 1
        else:
            created_count += 1
            new_product = Product(
                sku=sku,
                name=str(item.get("name") or sku).strip(),
                category_id=category_obj.id if category_obj else None,
                base_unit=str(item.get("unit") or "Cái").strip(),
                base_price=float(item.get("sale_price") or 0.0),
                cost_price=float(item.get("cost_price") or 0.0),
                min_stock=int(item.get("stock_available") or 0),
                manage_by_lot=False,
                status="ACTIVE",
            )
            db.add(new_product)
            db.flush()

            if new_product.base_price > 0:
                record_price_history(
                    db=db,
                    product_id=new_product.id,
                    old_price=None,
                    new_price=new_product.base_price,
                    changed_by=changed_by or "System",
                    note="Khởi tạo giá sản phẩm mới qua import",
                    reason="Khởi tạo giá sản phẩm mới qua import",
                )

            base_unit_entry = ProductUnit(
                product_id=new_product.id,
                unit_name=new_product.base_unit,
                conversion_rate=1.0,
            )
            db.add(base_unit_entry)

        # Sync with MOCK_PRODUCTS
        if sku in MOCK_PRODUCTS:
            if item.get("name"):
                MOCK_PRODUCTS[sku]["name"] = str(item["name"]).strip()
            if category_name:
                MOCK_PRODUCTS[sku]["category"] = category_name
            if item.get("sale_price") is not None:
                MOCK_PRODUCTS[sku]["sale_price"] = item["sale_price"]
            if item.get("cost_price") is not None:
                MOCK_PRODUCTS[sku]["cost_price"] = item["cost_price"]
            if item.get("stock_available") is not None:
                MOCK_PRODUCTS[sku]["stock_available"] = item["stock_available"]
            if item.get("unit"):
                MOCK_PRODUCTS[sku]["unit"] = item["unit"]
        else:
            MOCK_PRODUCTS[sku] = {
                "sku": sku,
                "name": str(item.get("name") or sku).strip(),
                "category": category_name,
                "sale_price": item.get("sale_price", 0),
                "cost_price": item.get("cost_price", 0),
                "stock_available": item.get("stock_available", 0),
                "unit": str(item.get("unit") or "Cái").strip(),
            }

    db.commit()
    return {
        "success": True,
        "message": f"Nhập danh mục sản phẩm thành công. Đã tạo mới {created_count}, cập nhật {updated_count} sản phẩm.",
        "created_count": created_count,
        "updated_count": updated_count,
        "total": created_count + updated_count,
    }


@router.get("/import/template")
def download_import_template(
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_IMPORT)),
) -> StreamingResponse:
    """Download standard Excel import template."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Products"
    worksheet.append(list(IMPORT_COLUMNS))
    worksheet.append(["SKU-001", "Nước ngọt Coca-Cola 330ml", "Nước giải khát", 10000, 7500, 100, "Lon"])
    worksheet.append(["SKU-002", "Bánh Cosy Mè 288g", "Bánh kẹo", 35000, 26000, 50, "Hộp"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="product-import-template.xlsx"'},
    )


@router.post("/import/preview")
async def preview_product_import(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_IMPORT)),
) -> dict[str, Any]:
    """Parse Excel file and return validation & preview list."""
    content = await _read_import_file(file)
    parsed_rows = _parse_import_workbook(content)
    existing_skus = _get_existing_skus(db)
    rows = _annotate_import_rows(parsed_rows, existing_skus)

    session_id = str(uuid.uuid4())
    IMPORT_SESSIONS[session_id] = rows

    return {
        "success": True,
        "session_id": session_id,
        "batch_id": session_id,
        "total": len(rows),
        "create_count": sum(1 for r in rows if r["action_type"] == "CREATE"),
        "update_count": sum(1 for r in rows if r["action_type"] == "UPDATE"),
        "error_count": sum(1 for r in rows if r["action_type"] == "ERROR"),
        "items": rows,
        "rows": rows,
    }


@router.post("/import/commit")
async def commit_product_import(
    request: Request,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_IMPORT)),
) -> dict[str, Any]:
    """Commit validated products into the database."""
    content_type = request.headers.get("content-type", "")

    if "multipart/form-data" in content_type:
        form = await request.form()
        file = form.get("file")
        if not file or not hasattr(file, "read"):
            raise HTTPException(status_code=400, detail="Thiếu tệp tải lên (file).")
        content = await _read_import_file(file)
        parsed_rows = _parse_import_workbook(content)
        existing_skus = _get_existing_skus(db)
        rows = _annotate_import_rows(parsed_rows, existing_skus)
        if any(r["action_type"] == "ERROR" for r in rows):
            raise HTTPException(
                status_code=422,
                detail={"message": "Không có dữ liệu nào được lưu vì tệp còn dòng lỗi.", "rows": rows},
            )
        return _apply_import_rows(rows, db, changed_by=user.username)

    try:
        body = await request.json()
    except Exception:
        body = {}

    session_id = body.get("session_id")
    items_to_commit: list[dict[str, Any]] = []

    if session_id and session_id in IMPORT_SESSIONS:
        cached_rows = IMPORT_SESSIONS.pop(session_id)
        if any(r.get("action_type") == "ERROR" for r in cached_rows):
            raise HTTPException(
                status_code=422,
                detail="Không thể lưu vì phiên nhập chứa dòng dữ liệu lỗi.",
            )
        items_to_commit = cached_rows
    elif "items" in body and isinstance(body["items"], list):
        items_to_commit = body["items"]
    else:
        raise HTTPException(
            status_code=400,
            detail="Thiếu dữ liệu session_id hoặc danh sách items để lưu.",
        )

    return _apply_import_rows(items_to_commit, db, changed_by=user.username)


@router.post("/import-excel")
async def import_excel_direct(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_IMPORT)),
) -> dict[str, Any]:
    """Direct Excel import endpoint matching API_SPECIFICATION.md POST /products/import-excel."""
    content = await _read_import_file(file)
    parsed_rows = _parse_import_workbook(content)
    existing_skus = _get_existing_skus(db)
    rows = _annotate_import_rows(parsed_rows, existing_skus)
    if any(r["action_type"] == "ERROR" for r in rows):
        raise HTTPException(
            status_code=422,
            detail={"message": "Không thể import vì tệp chứa dòng lỗi.", "rows": rows},
        )
    return _apply_import_rows(rows, db, changed_by=user.username)



@router.get("")
def list_products(
    keyword: Optional[str] = None,
    category_id: Optional[int] = Query(None, alias="categoryId"),
    category_id_legacy: Optional[int] = Query(None, alias="category_id"),
    status_filter: Optional[str] = Query(None, alias="status"),
    page: Optional[int] = None,
    limit: Optional[int] = None,
    skip: Optional[int] = None,
    offset: Optional[int] = None,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Get list of products with optional server-side filtering and pagination."""
    effective_cat_id = category_id if category_id is not None else category_id_legacy
    effective_offset = offset if offset is not None else skip
    
    if db.query(Product).count() > 0:
        query = db.query(Product)
        if keyword:
            query = query.filter(
                (Product.name.ilike(f"%{keyword}%")) | (Product.sku.ilike(f"%{keyword}%"))
            )
        if effective_cat_id is not None:
            query = query.filter(Product.category_id == effective_cat_id)
        if status_filter:
            query = query.filter(Product.status == status_filter)
            
        total_records = query.count()
        if effective_offset is not None and limit is not None:
            products = query.offset(effective_offset).limit(limit).all()
        elif page is not None and limit is not None:
            products = query.offset((page - 1) * limit).limit(limit).all()
        else:
            products = query.all()
        
        categories_map = {c.id: c.name for c in db.query(Category).all()}
        products_list = []
        for p in products:
            p_dict = {
                "id": p.id,
                "sku": p.sku,
                "name": p.name,
                "category_id": p.category_id,
                "categoryId": p.category_id,
                "category": categories_map.get(p.category_id) or (p.category.name if getattr(p, "category", None) else None),
                "baseUnit": p.base_unit,
                "unit": p.base_unit,
                "manageByLot": p.manage_by_lot,
                "minStock": p.min_stock,
                "basePrice": p.base_price,
                "sale_price": p.base_price,
                "costPrice": p.cost_price,
                "cost_price": p.cost_price,
                "margin": f"{round(((p.base_price - p.cost_price) / p.base_price * 100), 1)}%"
                if p.base_price and p.cost_price
                else None,
                "status": p.status,
                "imageUrl": p.image_url,
            }
            products_list.append(p_dict)
    else:
        mock_items = list(MOCK_PRODUCTS.values())
        if effective_cat_id is not None:
            mock_items = [p for p in mock_items if p.get("category_id") == effective_cat_id]
        if keyword:
            kw = keyword.lower()
            mock_items = [p for p in mock_items if kw in p.get("name", "").lower() or kw in p.get("sku", "").lower()]
        total_records = len(mock_items)
        if effective_offset is not None and limit is not None:
            mock_items = mock_items[effective_offset : effective_offset + limit]
        elif page is not None and limit is not None:
            mock_items = mock_items[(page - 1) * limit : page * limit]
        products_list = [p.copy() for p in mock_items]

    sanitized = sanitize_financial_data(products_list, user)
    
    if (effective_offset is not None and limit is not None) or (page is not None and limit is not None):
        res_limit = limit if limit is not None else 20
        res_page = page if page is not None else (effective_offset // res_limit + 1 if res_limit > 0 else 1)
        total_pages = (total_records + res_limit - 1) // res_limit if res_limit > 0 else 1
        return {
            "success": True,
            "code": 200,
            "message": "Thành công",
            "data": sanitized,
            "pagination": {
                "page": res_page,
                "limit": res_limit,
                "totalRecords": total_records,
                "totalPages": total_pages,
            },
        }
    return sanitized


@router.get("/{id_or_sku}")
def get_product_detail(
    id_or_sku: str,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Get product detail by ID or SKU."""
    product = None
    if id_or_sku.isdigit():
        product = db.query(Product).filter(Product.id == int(id_or_sku)).first()
    if not product:
        product = db.query(Product).filter(
            (Product.sku == id_or_sku) | (Product.sku == id_or_sku.upper())
        ).first()

    if product:
        units = db.query(ProductUnit).filter(ProductUnit.product_id == product.id).all()
        p_dict = {
            "id": product.id,
            "sku": product.sku,
            "name": product.name,
            "category_id": product.category_id,
            "categoryId": product.category_id,
            "category": product.category.name if product.category else None,
            "baseUnit": product.base_unit,
            "unit": product.base_unit,
            "manageByLot": product.manage_by_lot,
            "minStock": product.min_stock,
            "basePrice": product.base_price,
            "sale_price": product.base_price,
            "costPrice": product.cost_price,
            "cost_price": product.cost_price,
            "margin": f"{round(((product.base_price - product.cost_price) / product.base_price * 100), 1)}%"
            if product.base_price and product.cost_price
            else None,
            "status": product.status,
            "imageUrl": product.image_url,
            "units": [
                {
                    "id": u.id,
                    "unitName": u.unit_name,
                    "conversionRate": u.conversion_rate,
                    "barcode": u.barcode,
                }
                for u in units
            ],
        }
        return sanitize_financial_data(p_dict, user)

    mock_p = MOCK_PRODUCTS.get(id_or_sku.upper()) or MOCK_PRODUCTS.get(id_or_sku)
    if mock_p:
        return sanitize_financial_data(mock_p.copy(), user)

    raise HTTPException(status_code=404, detail=f"Sản phẩm '{id_or_sku}' không tồn tại")


@router.get("/{product_id}/units")
def list_product_units(
    product_id: int = PathParam(gt=0),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
    db: Session = Depends(get_db),
) -> list[dict[str, Any]]:
    """List configured conversion units for one product."""
    rows = (
        db.query(ProductUnit)
        .filter(ProductUnit.product_id == product_id)
        .order_by(ProductUnit.is_base_unit.desc(), ProductUnit.unit_name.asc())
        .all()
    )
    return [
        {
            "id": row.id,
            "product_id": row.product_id,
            "unit_name": row.unit_name,
            "conversion_rate": row.conversion_rate,
            "is_base_unit": row.is_base_unit,
        }
        for row in rows
    ]


@router.post("/{product_id}/units")
def upsert_product_unit(
    payload: ProductUnitRequest,
    product_id: int = PathParam(gt=0),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_UNIT_MANAGE)),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Create or update a unit; at most one unit per product can be the base."""
    conversion_rate = 1.0 if payload.is_base_unit else payload.conversion_rate
    try:
        row = (
            db.query(ProductUnit)
            .filter(
                ProductUnit.product_id == product_id,
                ProductUnit.unit_name == payload.unit_name,
            )
            .first()
        )
        if row is None:
            row = ProductUnit(
                product_id=product_id,
                unit_name=payload.unit_name,
                conversion_rate=conversion_rate,
                is_base_unit=payload.is_base_unit,
            )
            db.add(row)
        else:
            row.conversion_rate = conversion_rate
            row.is_base_unit = payload.is_base_unit

        if payload.is_base_unit:
            (
                db.query(ProductUnit)
                .filter(
                    ProductUnit.product_id == product_id,
                    ProductUnit.unit_name != payload.unit_name,
                )
                .update({ProductUnit.is_base_unit: False}, synchronize_session=False)
            )
        db.commit()
        db.refresh(row)
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không thể lưu đơn vị tính sản phẩm.",
        ) from exc

    return {
        "id": row.id,
        "product_id": row.product_id,
        "unit_name": row.unit_name,
        "conversion_rate": row.conversion_rate,
        "is_base_unit": row.is_base_unit,
    }



class MoveProductCategoryRequest(BaseModel):
    target_category_id: Optional[int] = Field(default=None, gt=0, description="ID nhóm hàng mới cần chuyển đến")
    category_id: Optional[int] = Field(default=None, gt=0)
    categoryId: Optional[int] = Field(default=None, gt=0)

    @model_validator(mode="before")
    @classmethod
    def resolve_target_category_id(cls, data: Any) -> Any:
        if isinstance(data, dict):
            target_id = data.get("target_category_id") or data.get("category_id") or data.get("categoryId")
            if target_id is not None:
                data["target_category_id"] = target_id
        return data




@router.patch("/{sku}/category")
def move_product_category(
    sku: str,
    payload: MoveProductCategoryRequest,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Move a single product to a new target category (SCRUM-76 AC 2)."""
    if not payload.target_category_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Vui lòng cung cấp target_category_id hoặc category_id hợp lệ.",
        )

    from app.models.models import RoleCode
    if user.role not in (RoleCode.SALES_MANAGER.value, RoleCode.ADMIN.value):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Chỉ Quản lý kinh doanh và Quản trị viên mới có quyền chuyển nhóm hàng sản phẩm.",
        )

    from app.api.v1.endpoints.product_categories import get_category_by_id
    target_cat = get_category_by_id(payload.target_category_id)
    if not target_cat:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Nhóm hàng đích với ID '{payload.target_category_id}' không tồn tại.",
        )

    # Update in DB if exists
    product = db.query(Product).filter((Product.sku == sku) | (Product.sku == sku.upper())).first()
    old_category_id = None
    old_category_name = None
    product_name = sku
    if product:
        old_category_id = product.category_id
        product.category_id = target_cat["id"]
        product_name = product.name
        db.commit()
        db.refresh(product)

    # Also update in MOCK_PRODUCTS if present
    mock_p = MOCK_PRODUCTS.get(sku.upper()) or MOCK_PRODUCTS.get(sku)
    if mock_p:
        old_category_id = old_category_id or mock_p.get("category_id")
        old_category_name = mock_p.get("category")
        product_name = mock_p.get("name", product_name)
        mock_p["category_id"] = target_cat["id"]
        mock_p["category"] = target_cat["name"]

    if not product and not mock_p:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Sản phẩm với mã SKU '{sku}' không tồn tại.",
        )

    return {
        "message": f"Chuyển sản phẩm '{sku}' sang nhóm '{target_cat['name']}' thành công.",
        "sku": sku,
        "name": product_name,
        "old_category_id": old_category_id,
        "old_category_name": old_category_name,
        "new_category_id": target_cat["id"],
        "new_category_name": target_cat["name"],
    }


class UpdateProductPriceRequest(BaseModel):
    price: float = Field(..., ge=0, description="Giá bán mới của sản phẩm")
    note: Optional[str] = None
    reason: Optional[str] = None
    effective_date: Optional[Union[datetime, str]] = None

    @model_validator(mode="before")
    @classmethod
    def extract_fields(cls, data: Any) -> Any:
        if isinstance(data, dict):
            p = data.get("price")
            if p is None:
                p = (
                    data.get("new_price")
                    or data.get("newPrice")
                    or data.get("sale_price")
                    or data.get("base_price")
                )
            if p is not None:
                data["price"] = float(p)
            note_val = data.get("note") or data.get("reason")
            if note_val:
                data["note"] = note_val
                data["reason"] = note_val
            eff_val = (
                data.get("effective_date")
                or data.get("effectiveDate")
                or data.get("changed_at")
                or data.get("changedAt")
            )
            if eff_val:
                data["effective_date"] = eff_val
        return data


UpdateProductPriceRequest.model_rebuild()


def _resolve_product(db: Session, id_or_sku: Union[int, str]) -> Optional[Product]:
    id_str = str(id_or_sku).strip()
    product = None
    if id_str.isdigit():
        product = db.query(Product).filter(Product.id == int(id_str)).first()
    if not product:
        product = db.query(Product).filter(
            (Product.sku == id_str) | (Product.sku == id_str.upper())
        ).first()
    if not product:
        mock_p = MOCK_PRODUCTS.get(id_str.upper()) or MOCK_PRODUCTS.get(id_str)
        if mock_p:
            sku = mock_p.get("sku", id_str.upper())
            product = db.query(Product).filter(Product.sku == sku).first()
            if not product:
                product = Product(
                    sku=sku,
                    name=mock_p.get("name", sku),
                    base_unit=mock_p.get("unit", "Cái"),
                    base_price=float(mock_p.get("sale_price", 0.0)),
                    cost_price=float(mock_p.get("cost_price", 0.0)),
                    category_id=mock_p.get("category_id"),
                    status="ACTIVE",
                )
                db.add(product)
                db.flush()
    return product


# SCRUM-18: Price history endpoint
@router.get("/{id_or_sku}/price-history")
def get_product_price_history_endpoint(
    id_or_sku: str,
    fromDate: Optional[str] = Query(None, alias="fromDate"),
    toDate: Optional[str] = Query(None, alias="toDate"),
    from_date: Optional[str] = Query(None),
    to_date: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_roles("Sales Manager", "Admin")),
) -> Any:
    """Retrieve product price change history (Story S3-02 / SCRUM-18)."""
    product = _resolve_product(db, id_or_sku)
    if not product:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Sản phẩm '{id_or_sku}' không tồn tại.",
        )
    effective_from = fromDate or from_date
    effective_to = toDate or to_date
    histories = get_price_history(db, product.id, from_date=effective_from, to_date=effective_to)
    return [serialize_price_history(h) for h in histories]


# SCRUM-18: Price update endpoint
@router.put("/{id_or_sku}/price", operation_id="update_product_price_put")
@router.patch("/{id_or_sku}/price", operation_id="update_product_price_patch")
def update_product_price_endpoint(
    id_or_sku: str,
    payload: UpdateProductPriceRequest,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_roles("Sales Manager", "Admin")),
) -> Any:
    """Update product price and automatically log price history audit record (SCRUM-18)."""
    product = _resolve_product(db, id_or_sku)
    if not product:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Sản phẩm '{id_or_sku}' không tồn tại.",
        )

    old_price = float(product.base_price) if product.base_price is not None else 0.0
    new_price = float(payload.price)
    effective_date = parse_date_boundary(payload.effective_date) if payload.effective_date else None

    # Update product price
    product._suppress_auto_price_history = True
    product.base_price = new_price
    if product.sku in MOCK_PRODUCTS:
        MOCK_PRODUCTS[product.sku]["sale_price"] = new_price

    note = payload.note or payload.reason or "Cập nhật giá bán sản phẩm"
    history = record_price_history(
        db=db,
        product_id=product.id,
        old_price=old_price,
        new_price=new_price,
        changed_by=user.username or user.role,
        note=note,
        reason=note,
        effective_date=effective_date,
    )
    db.commit()
    db.refresh(product)
    db.refresh(history)

    return {
        "success": True,
        "message": f"Cập nhật giá bán sản phẩm '{product.sku}' thành công.",
        "product_id": product.id,
        "sku": product.sku,
        "old_price": old_price,
        "new_price": new_price,
        "history": serialize_price_history(history),
    }


# SCRUM-18: Reject any modification or deletion of price history (Audit Log immutability)
@router.api_route("/{id_or_sku}/price-history", methods=["POST", "PUT", "PATCH", "DELETE"], include_in_schema=False)
@router.api_route("/{id_or_sku}/price-history/{history_id}", methods=["POST", "PUT", "PATCH", "DELETE"], include_in_schema=False)
def reject_price_history_mutation(id_or_sku: str, history_id: Optional[str] = None) -> Any:
    """Audit Log immutability enforcement: Price history is strictly read-only."""
    raise HTTPException(
        status_code=status.HTTP_405_METHOD_NOT_ALLOWED,
        detail="Lịch sử thay đổi giá là nhật ký kiểm toán bất biến (Audit Log), không cho phép chỉnh sửa hoặc xóa.",
    )


# Compat router for /products/... without /api/v1 prefix
compat_router = APIRouter(prefix="/products", tags=["Products Compat"])
compat_router.add_api_route(
    "/{id_or_sku}/price-history",
    get_product_price_history_endpoint,
    methods=["GET"],
    include_in_schema=False,
)
compat_router.add_api_route(
    "/{id_or_sku}/price",
    update_product_price_endpoint,
    methods=["PUT", "PATCH"],
    include_in_schema=False,
)
compat_router.add_api_route(
    "/{id_or_sku}/price-history",
    reject_price_history_mutation,
    methods=["POST", "PUT", "PATCH", "DELETE"],
    include_in_schema=False,
)
compat_router.add_api_route(
    "/{id_or_sku}/price-history/{history_id}",
    reject_price_history_mutation,
    methods=["POST", "PUT", "PATCH", "DELETE"],
    include_in_schema=False,
)


def _check_product_transactions(product_id: int, db: Session) -> tuple[bool, str]:
    """Check if a product has been used in any transactions (orders, receipts, transfers, audits, dispatches, returns, adjustments)."""
    if db.query(OrderItem).filter(OrderItem.product_id == product_id).first():
        return True, "đơn đặt hàng"
    if db.query(WarehouseReceiptItem).filter(WarehouseReceiptItem.product_id == product_id).first():
        return True, "phiếu nhập kho"
    if db.query(WarehouseTransferItem).filter(WarehouseTransferItem.product_id == product_id).first():
        return True, "phiếu điều chuyển kho"
    if db.query(WarehouseAuditItem).filter(WarehouseAuditItem.product_id == product_id).first():
        return True, "phiếu kiểm kê kho"
    if db.query(DispatchItem).filter(DispatchItem.product_id == product_id).first():
        return True, "phiếu xuất kho"
    if db.query(ReturnOrderItem).filter(ReturnOrderItem.product_id == product_id).first():
        return True, "phiếu trả hàng"
    if db.query(InventoryAdjustmentItem).filter(InventoryAdjustmentItem.product_id == product_id).first():
        return True, "phiếu điều chỉnh tồn kho"
    return False, ""


@router.post("", status_code=status.HTTP_201_CREATED)
def create_product(
    payload: ProductCreateRequest,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Create a new product (Sales Manager only)."""
    if user.role not in (RoleCode.SALES_MANAGER.value, RoleCode.ADMIN.value):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Chỉ Quản lý kinh doanh mới có quyền khai báo danh mục sản phẩm.",
        )

    sku_clean = payload.sku.strip().upper()
    existing = db.query(Product).filter(Product.sku == sku_clean).first()
    if existing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Mã SKU '{sku_clean}' đã tồn tại trong hệ thống. Vui lòng chọn mã khác.",
        )

    if payload.categoryId:
        cat = db.query(Category).filter(Category.id == payload.categoryId).first()
        if not cat:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Không tìm thấy nhóm hàng có ID {payload.categoryId}.",
            )

    cost_price = payload.costPrice or 0.0
    if not can_view_financials(user.role):
        cost_price = 0.0

    product = Product(
        sku=sku_clean,
        name=payload.name.strip(),
        category_id=payload.categoryId,
        base_unit=payload.baseUnit.strip(),
        manage_by_lot=payload.manageByLot,
        min_stock=payload.minStock,
        base_price=payload.basePrice or 0.0,
        cost_price=cost_price,
        image_url=payload.imageUrl,
        status=payload.status or "ACTIVE",
    )
    db.add(product)
    db.flush()

    # Base unit entry
    base_unit_entry = ProductUnit(
        product_id=product.id,
        unit_name=product.base_unit,
        conversion_rate=1.0,
        is_base_unit=True,
    )
    db.add(base_unit_entry)

    # Additional units if provided
    if payload.units:
        for u in payload.units:
            if u.unitName.strip() != product.base_unit:
                conv_rate = u.conversionRate if u.conversionRate > 0 else 1.0
                extra_unit = ProductUnit(
                    product_id=product.id,
                    unit_name=u.unitName.strip(),
                    conversion_rate=conv_rate,
                    barcode=u.barcode,
                    is_base_unit=False,
                )
                db.add(extra_unit)

    db.commit()
    db.refresh(product)

    res_data = {
        "id": product.id,
        "sku": product.sku,
        "name": product.name,
        "categoryId": product.category_id,
        "category_id": product.category_id,
        "baseUnit": product.base_unit,
        "manageByLot": product.manage_by_lot,
        "minStock": product.min_stock,
        "basePrice": product.base_price,
        "costPrice": product.cost_price,
        "imageUrl": product.image_url,
        "status": product.status,
    }
    return {
        "success": True,
        "code": 201,
        "message": f"Tạo sản phẩm '{product.name}' thành công.",
        "data": sanitize_financial_data(res_data, user),
    }


@router.put("/{product_id}")
def update_product(
    product_id: int,
    payload: ProductUpdateRequest,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Update product information and cost price (Sales Manager only)."""
    if user.role not in (RoleCode.SALES_MANAGER.value, RoleCode.ADMIN.value):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Chỉ Quản lý kinh doanh mới có quyền chỉnh sửa thông tin sản phẩm.",
        )

    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Không tìm thấy sản phẩm có ID {product_id}.",
        )

    if payload.name is not None:
        product.name = payload.name.strip()
    if payload.categoryId is not None:
        if payload.categoryId > 0:
            cat = db.query(Category).filter(Category.id == payload.categoryId).first()
            if not cat:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Không tìm thấy nhóm hàng có ID {payload.categoryId}.",
                )
            product.category_id = payload.categoryId
        else:
            product.category_id = None
    if payload.baseUnit is not None:
        product.base_unit = payload.baseUnit.strip()
        base_unit_rec = db.query(ProductUnit).filter(ProductUnit.product_id == product.id, ProductUnit.is_base_unit.is_(True)).first()
        if base_unit_rec:
            base_unit_rec.unit_name = product.base_unit
    if payload.manageByLot is not None:
        product.manage_by_lot = payload.manageByLot
    if payload.minStock is not None:
        product.min_stock = payload.minStock
    if payload.basePrice is not None:
        product.base_price = payload.basePrice
    if payload.costPrice is not None:
        if not can_view_financials(user.role):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Chỉ Quản lý kinh doanh mới có quyền xem và sửa giá vốn.",
            )
        product.cost_price = payload.costPrice
    if payload.imageUrl is not None:
        product.image_url = payload.imageUrl
    if payload.status is not None:
        product.status = payload.status

    if payload.units is not None:
        db.query(ProductUnit).filter(ProductUnit.product_id == product.id, ProductUnit.is_base_unit.is_(False)).delete()
        for u in payload.units:
            if u.unitName.strip() != product.base_unit:
                db.add(ProductUnit(
                    product_id=product.id,
                    unit_name=u.unitName.strip(),
                    conversion_rate=u.conversionRate if u.conversionRate > 0 else 1.0,
                    barcode=u.barcode,
                    is_base_unit=False,
                ))

    db.commit()
    db.refresh(product)

    res_data = {
        "id": product.id,
        "sku": product.sku,
        "name": product.name,
        "categoryId": product.category_id,
        "category_id": product.category_id,
        "baseUnit": product.base_unit,
        "manageByLot": product.manage_by_lot,
        "minStock": product.min_stock,
        "basePrice": product.base_price,
        "costPrice": product.cost_price,
        "imageUrl": product.image_url,
        "status": product.status,
    }
    return {
        "success": True,
        "code": 200,
        "message": f"Cập nhật sản phẩm '{product.name}' thành công.",
        "data": sanitize_financial_data(res_data, user),
    }


@router.patch("/{product_id}/status")
def toggle_product_status(
    product_id: int,
    payload: ProductStatusUpdateRequest,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Toggle product status (ACTIVE or INACTIVE - ngừng kinh doanh)."""
    if user.role not in (RoleCode.SALES_MANAGER.value, RoleCode.ADMIN.value):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Chỉ Quản lý kinh doanh mới có quyền chuyển trạng thái sản phẩm.",
        )

    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Không tìm thấy sản phẩm có ID {product_id}.",
        )

    product.status = payload.status
    db.commit()
    db.refresh(product)

    status_label = "Đang kinh doanh" if product.status == "ACTIVE" else "Ngừng kinh doanh"
    return {
        "success": True,
        "code": 200,
        "message": f"Đã chuyển trạng thái sản phẩm '{product.name}' sang: {status_label}.",
        "data": {
            "id": product.id,
            "sku": product.sku,
            "status": product.status,
        },
    }


@router.delete("/{product_id}")
def delete_product(
    product_id: int,
    db: Session = Depends(get_db),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Delete a product. If transaction exists, forbid deletion and advise setting to INACTIVE."""
    if user.role not in (RoleCode.SALES_MANAGER.value, RoleCode.ADMIN.value):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Chỉ Quản lý kinh doanh mới có quyền xoá sản phẩm.",
        )

    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Không tìm thấy sản phẩm có ID {product_id}.",
        )

    has_tx, tx_name = _check_product_transactions(product.id, db)
    if has_tx:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Sản phẩm '{product.sku}' đã phát sinh giao dịch ({tx_name}) nên không thể xoá, chỉ có thể chuyển sang 'Ngừng kinh doanh'.",
        )

    db.query(ProductUnit).filter(ProductUnit.product_id == product.id).delete()
    db.delete(product)
    db.commit()

    return {
        "success": True,
        "code": 200,
        "message": f"Đã xoá hoàn toàn sản phẩm '{product.sku}' khỏi hệ thống.",
        "data": {"id": product_id, "sku": product.sku},
    }



