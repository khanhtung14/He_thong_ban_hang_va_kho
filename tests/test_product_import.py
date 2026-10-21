import copy
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from src.backend import products
from src.backend.main import app
from src.backend.rbac import create_access_token


client = TestClient(app)


def auth_headers(role="Sales Manager"):
    token = create_access_token({"sub": "import_user", "role": role})
    return {"Authorization": f"Bearer {token}"}


def xlsx_file(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(products.IMPORT_COLUMNS)
    for row in rows:
        worksheet.append(row)
    content = io.BytesIO()
    workbook.save(content)
    return content.getvalue()


def test_import_template_is_downloadable_xlsx():
    response = client.get("/api/v1/products/import/template", headers=auth_headers())

    assert response.status_code == 200
    assert "product-import-template.xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    assert tuple(next(workbook.active.iter_rows(values_only=True))) == products.IMPORT_COLUMNS
    workbook.close()


def test_preview_marks_existing_skus_and_reports_row_errors(monkeypatch):
    monkeypatch.setattr(products, "MOCK_PRODUCTS", copy.deepcopy(products.MOCK_PRODUCTS))
    content = xlsx_file(
        [
            ["SKU-001", "Updated", "Beverages", 510000, None, None, None],
            ["SKU-003", "New", "Snacks", None, None, 5, "Box"],
            ["SKU-004", "Missing category", "", None, None, None, None],
        ]
    )

    response = client.post(
        "/api/v1/products/import/preview",
        headers=auth_headers(),
        files={"file": ("products.xlsx", content)},
    )

    assert response.status_code == 200
    data = response.json()
    assert (data["create_count"], data["update_count"], data["error_count"]) == (1, 1, 1)
    assert data["rows"][0]["action"] == "update"
    assert "tồn tại" in data["rows"][0]["message"]
    assert data["rows"][1]["action"] == "create"
    assert data["rows"][2]["action"] == "error"
    assert "Danh mục" in data["rows"][2]["message"]
    assert data["rows"][2]["row"] == 4


def test_import_commit_updates_existing_sku_and_creates_new_sku(monkeypatch):
    monkeypatch.setattr(products, "MOCK_PRODUCTS", copy.deepcopy(products.MOCK_PRODUCTS))
    content = xlsx_file(
        [
            ["SKU-001", "Updated name", "New category", 520000, None, None, None],
            ["SKU-003", "New product", "Snacks", 30000, None, 12, "Box"],
        ]
    )

    response = client.post(
        "/api/v1/products/import/commit",
        headers=auth_headers(),
        files={"file": ("products.xlsx", content)},
    )

    assert response.status_code == 200
    assert response.json()["created_count"] == 1
    assert response.json()["updated_count"] == 1
    assert products.MOCK_PRODUCTS["SKU-001"]["name"] == "Updated name"
    assert products.MOCK_PRODUCTS["SKU-001"]["category"] == "New category"
    assert products.MOCK_PRODUCTS["SKU-003"]["name"] == "New product"


def test_import_commit_rejects_invalid_batch_without_partial_updates(monkeypatch):
    monkeypatch.setattr(products, "MOCK_PRODUCTS", copy.deepcopy(products.MOCK_PRODUCTS))
    original_name = products.MOCK_PRODUCTS["SKU-001"]["name"]
    content = xlsx_file(
        [
            ["SKU-001", "Should not apply", "Beverages", None, None, None, None],
            ["SKU-003", "Invalid", "", None, None, None, None],
        ]
    )

    response = client.post(
        "/api/v1/products/import/commit",
        headers=auth_headers(),
        files={"file": ("products.xlsx", content)},
    )

    assert response.status_code == 422
    assert products.MOCK_PRODUCTS["SKU-001"]["name"] == original_name
    assert "rows" in response.json()["detail"]


def test_only_product_import_roles_can_upload_files():
    response = client.post(
        "/api/v1/products/import/preview",
        headers=auth_headers("Sales Rep"),
        files={"file": ("products.xlsx", xlsx_file([["SKU-003", "New", "Beverages", None, None, None, None]]))},
    )

    assert response.status_code == 403


def test_preview_rejects_corrupt_xlsx():
    response = client.post(
        "/api/v1/products/import/preview",
        headers=auth_headers(),
        files={"file": ("products.xlsx", b"not an xlsx file")},
    )

    assert response.status_code == 422


def test_preview_annotates_new_and_existing_categories():
    content = xlsx_file(
        [
            ["SKU-NEW-01", "Trà Xanh 0 Độ", "Nước giải khát", 12000, 8000, 50, "Chai"],
            ["SKU-NEW-02", "Bánh Mì Tươi", "Bánh Tươi Đặc Biệt", 15000, 10000, 20, "Gói"],
        ]
    )

    response = client.post(
        "/api/v1/products/import/preview",
        headers=auth_headers(),
        files={"file": ("products.xlsx", content)},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 2
    # Row 1 has existing category "Nước giải khát"
    row1 = data["rows"][0]
    assert row1["category"] == "Nước giải khát"
    assert row1["is_new_category"] is False
    assert row1["category_status"] == "EXISTING"

    # Row 2 has non-existent category "Bánh Tươi Đặc Biệt"
    row2 = data["rows"][1]
    assert row2["category"] == "Bánh Tươi Đặc Biệt"
    assert row2["is_new_category"] is True
    assert row2["category_status"] == "NEW"
    assert "Danh mục mới (sẽ được tự động tạo)" in row2["category_note"]
    assert "Danh mục mới (sẽ được tự động tạo)" in row2["message"]
    # Row 2 should NOT be marked as an error!
    assert row2["action_type"] == "CREATE"


def test_commit_find_or_create_category_with_cache_and_linkage():
    from app.core.database import SessionLocal
    from app.models.models import Category, Product

    unique_cat_name = "Thực phẩm đóng hộp Hảo Hạng"
    content = xlsx_file(
        [
            # Two products sharing the same new category name to test cache
            ["SKU-AUTO-CAT-01", "Cá hộp Ba Cô Gái", f"  {unique_cat_name}  ", 25000, 18000, 100, "Hộp"],
            ["SKU-AUTO-CAT-02", "Pate Gan Heo Hạ Long", unique_cat_name.lower(), 30000, 22000, 80, "Hộp"],
            # One product with existing category in different case
            ["SKU-AUTO-CAT-03", "Coca Cola 330ml", "  nước giải khát  ", 10000, 7000, 50, "Lon"],
        ]
    )

    response = client.post(
        "/api/v1/products/import/commit",
        headers=auth_headers(),
        files={"file": ("products.xlsx", content)},
    )

    assert response.status_code == 200
    assert response.json()["success"] is True

    db = SessionLocal()
    try:
        # 1. Verify only ONE Category was created for unique_cat_name despite two rows
        created_cats = db.query(Category).filter(Category.name.ilike(unique_cat_name)).all()
        assert len(created_cats) == 1
        cat_obj = created_cats[0]
        assert cat_obj.code is not None
        assert len(cat_obj.code) > 0

        # 2. Verify products were created and linked to the new Category ID
        p1 = db.query(Product).filter(Product.sku == "SKU-AUTO-CAT-01").first()
        p2 = db.query(Product).filter(Product.sku == "SKU-AUTO-CAT-02").first()
        assert p1 is not None
        assert p2 is not None
        assert p1.category_id == cat_obj.id
        assert p2.category_id == cat_obj.id

        # 3. Verify case-insensitive existing category was correctly linked to existing category
        p3 = db.query(Product).filter(Product.sku == "SKU-AUTO-CAT-03").first()
        assert p3 is not None
        existing_cat = db.query(Category).filter(Category.name.ilike("Nước giải khát")).first()
        assert existing_cat is not None
        assert p3.category_id == existing_cat.id

        # 4. Verify category tree endpoint now returns this new category
        tree_res = client.get("/api/v1/categories/tree", headers=auth_headers())
        assert tree_res.status_code == 200
        tree = tree_res.json()

        def get_all_ids(nodes):
            res = []
            for n in nodes:
                res.append(n["id"])
                if n.get("children"):
                    res.extend(get_all_ids(n["children"]))
            return res

        cat_ids_in_tree = get_all_ids(tree)
        assert cat_obj.id in cat_ids_in_tree
    finally:
        # Cleanup
        db.query(Product).filter(Product.sku.in_(["SKU-AUTO-CAT-01", "SKU-AUTO-CAT-02", "SKU-AUTO-CAT-03"])).delete()
        db.query(Category).filter(Category.name == unique_cat_name).delete()
        db.commit()
        db.close()


def test_direct_import_excel_auto_creates_category():
    from app.core.database import SessionLocal
    from app.models.models import Category, Product

    unique_cat = "Dụng cụ nhà bếp tiện lợi"
    content = xlsx_file(
        [
            ["SKU-DIR-01", "Dao gọt hoa quả", unique_cat, 45000, 30000, 20, "Cây"],
        ]
    )

    response = client.post(
        "/api/v1/products/import-excel",
        headers=auth_headers(),
        files={"file": ("products.xlsx", content)},
    )

    assert response.status_code == 200
    db = SessionLocal()
    try:
        cat = db.query(Category).filter(Category.name == unique_cat).first()
        assert cat is not None
        prod = db.query(Product).filter(Product.sku == "SKU-DIR-01").first()
        assert prod is not None
        assert prod.category_id == cat.id
    finally:
        db.query(Product).filter(Product.sku == "SKU-DIR-01").delete()
        db.query(Category).filter(Category.name == unique_cat).delete()
        db.commit()
        db.close()


def test_import_excel_creates_inventory_record_at_default_warehouse():
    """S2-08: Kiểm tra import sản phẩm từ file Excel tự động tạo/cập nhật bản ghi trong bảng inventory gắn với Kho mặc định."""
    from app.core.database import SessionLocal
    from app.models.models import Category, Product, Inventory, Warehouse

    sku_test = "SKU-TEST-INV-150"
    cat_test = "Đồ uống kiểm thử tồn kho"
    content = xlsx_file(
        [
            [sku_test, "Sản phẩm test tồn kho", cat_test, 25000, 18000, 150, "Chai"],
        ]
    )

    response = client.post(
        "/api/v1/products/import-excel",
        headers=auth_headers(),
        files={"file": ("products.xlsx", content)},
    )
    assert response.status_code == 200

    db = SessionLocal()
    try:
        wh = db.query(Warehouse).filter(Warehouse.id == 1).first() or db.query(Warehouse).first()
        assert wh is not None

        prod = db.query(Product).filter(Product.sku == sku_test).first()
        assert prod is not None

        inv = db.query(Inventory).filter(
            Inventory.product_id == prod.id,
            Inventory.warehouse_id == wh.id,
        ).first()
        assert inv is not None
        assert inv.physical_qty == 150.0
        assert inv.quantity_on_hand == 150.0
        assert inv.reserved_qty == 0.0
        assert inv.reserved_quantity == 0.0
        assert inv.available_qty == 150.0
        assert inv.available_quantity == 150.0

        # Kiểm tra endpoint API available-stock
        res_stock = client.get(
            f"/api/v1/inventory/available-stock?sku={sku_test}&warehouseId={wh.id}",
            headers=auth_headers(),
        )
        assert res_stock.status_code == 200
        stock_data = res_stock.json()["data"]
        assert stock_data["physicalQty"] == 150.0
        assert stock_data["availableQty"] == 150.0
        assert stock_data["available_quantity"] == 150.0
    finally:
        prod_del = db.query(Product).filter(Product.sku == sku_test).first()
        if prod_del:
            db.query(Inventory).filter(Inventory.product_id == prod_del.id).delete()
        db.commit()
        db.close()


def test_import_excel_stock_quantity_alias_and_default_zero():
    """S2-08: Kiểm tra cột stock_quantity alias và nếu để trống tồn kho thì mặc định là 0."""
    from app.core.database import SessionLocal
    from app.models.models import Product, Inventory, Warehouse

    sku_alias = "SKU-TEST-ALIAS-01"
    sku_empty = "SKU-TEST-EMPTY-02"

    wb = Workbook()
    ws = wb.active
    # Sử dụng cột stock_quantity thay vì stock_available
    ws.append(["sku", "name", "category", "sale_price", "cost_price", "stock_quantity", "unit"])
    ws.append([sku_alias, "Sản phẩm dùng header stock_quantity", "Đồ uống", 15000, 10000, 75, "Chai"])
    ws.append([sku_empty, "Sản phẩm để trống tồn kho", "Đồ uống", 20000, 15000, None, "Lon"])
    content = io.BytesIO()
    wb.save(content)
    xlsx_bytes = content.getvalue()

    response = client.post(
        "/api/v1/products/import-excel",
        headers=auth_headers(),
        files={"file": ("products.xlsx", xlsx_bytes)},
    )
    assert response.status_code == 200

    db = SessionLocal()
    try:
        wh = db.query(Warehouse).filter(Warehouse.id == 1).first() or db.query(Warehouse).first()

        # Kiểm tra sản phẩm dùng header stock_quantity
        p_alias = db.query(Product).filter(Product.sku == sku_alias).first()
        assert p_alias is not None
        inv_alias = db.query(Inventory).filter(Inventory.product_id == p_alias.id, Inventory.warehouse_id == wh.id).first()
        assert inv_alias is not None
        assert inv_alias.quantity_on_hand == 75.0
        assert inv_alias.reserved_quantity == 0.0
        assert inv_alias.available_quantity == 75.0

        # Kiểm tra sản phẩm để trống tồn kho -> mặc định 0
        p_empty = db.query(Product).filter(Product.sku == sku_empty).first()
        assert p_empty is not None
        inv_empty = db.query(Inventory).filter(Inventory.product_id == p_empty.id, Inventory.warehouse_id == wh.id).first()
        assert inv_empty is not None
        assert inv_empty.quantity_on_hand == 0.0
        assert inv_empty.reserved_quantity == 0.0
        assert inv_empty.available_quantity == 0.0
    finally:
        for s in (sku_alias, sku_empty):
            p = db.query(Product).filter(Product.sku == s).first()
            if p:
                db.query(Inventory).filter(Inventory.product_id == p.id).delete()
        db.commit()
        db.close()

