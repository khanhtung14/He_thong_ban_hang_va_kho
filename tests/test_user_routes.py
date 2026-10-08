from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from src.backend.security import require_admin
from src.backend.user_routes import get_db, router


@pytest.fixture
def db_session_and_client():
    """Tạo database in-memory và FastAPI TestClient để kiểm thử user_routes."""
    test_app = FastAPI()
    test_app.include_router(router)

    # Sử dụng SQLite in-memory database với StaticPool để giữ nguyên bảng giữa các session
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    # Tạo bảng users và user_sessions tương tự schema MySQL
    with engine.connect() as conn:
        conn.execute(
            text(
                """
                CREATE TABLE users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username VARCHAR(100) NOT NULL UNIQUE,
                    is_active BOOLEAN NOT NULL DEFAULT 1
                );
                """
            )
        )
        conn.execute(
            text(
                """
                CREATE TABLE user_sessions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL,
                    revoked_at DATETIME
                );
                """
            )
        )
        # Thêm 2 user mẫu: ID 1 (sales01) và ID 2 (accountant01)
        conn.execute(text("INSERT INTO users (id, username, is_active) VALUES (1, 'sales01', 1);"))
        conn.execute(text("INSERT INTO users (id, username, is_active) VALUES (2, 'accountant01', 1);"))
        conn.execute(text("INSERT INTO user_sessions (id, user_id, revoked_at) VALUES (1, 1, NULL);"))
        conn.commit()

    def override_get_db():
        session = TestingSessionLocal()
        try:
            yield session
        finally:
            session.close()

    test_app.dependency_overrides[get_db] = override_get_db
    # These tests exercise route behavior as an authenticated administrator.
    # Production requests still use require_admin and must present a valid session.
    test_app.dependency_overrides[require_admin] = lambda: SimpleNamespace(
        id=999,
        roles=[SimpleNamespace(code="ADMIN")],
    )

    with TestClient(test_app) as client:
        yield client, engine


# =====================================================================
# 1. Test validation khi không có lý do hoặc lý do rỗng (Lỗi 400)
# =====================================================================
@pytest.mark.parametrize(
    "empty_reason",
    [
        "",
        "   ",
        "\t\n",
    ],
)
def test_lock_user_missing_or_blank_reason_returns_400(db_session_and_client, empty_reason):
    client, _ = db_session_and_client
    payload = {"user_id": "sales01", "reason": empty_reason}

    response = client.post("/api/users/lock", json=payload)

    assert response.status_code == 400
    assert response.json()["detail"] == "Bắt buộc phải nhập lý do khóa tài khoản."


def test_lock_user_blank_user_id_returns_400(db_session_and_client):
    client, _ = db_session_and_client
    payload = {"user_id": "   ", "reason": "Có lý do hợp lệ"}

    response = client.post("/api/users/lock", json=payload)

    assert response.status_code == 400
    assert response.json()["detail"] == "Vui lòng cung cấp ID hoặc username của người dùng."


# =====================================================================
# 2. Test khóa tài khoản thành công và xác nhận UPDATE is_active = False
# =====================================================================
def test_lock_user_success_updates_is_active_to_false(db_session_and_client):
    client, engine = db_session_and_client
    payload = {"user_id": "sales01", "reason": "Nhân viên nghỉ việc"}

    # Trước khi khóa: is_active = 1 (True)
    with engine.connect() as conn:
        status_before = conn.execute(text("SELECT is_active FROM users WHERE username = 'sales01'")).scalar()
        assert bool(status_before) is True

    # Gọi API khóa tài khoản
    response = client.post("/api/users/lock", json=payload)

    # Xác nhận response thành công và không bị 500
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "đã bị khóa" in data["message"]

    # Xác nhận trực tiếp trong Database: is_active đã được cập nhật thành False (0)
    with engine.connect() as conn:
        status_after = conn.execute(text("SELECT is_active FROM users WHERE username = 'sales01'")).scalar()
        assert bool(status_after) is False, f"Expected is_active=False, got {status_after}"

        # Xác nhận session của user cũng được thu hồi (revoked_at không còn NULL)
        revoked_at = conn.execute(text("SELECT revoked_at FROM user_sessions WHERE user_id = 1")).scalar()
        assert revoked_at is not None


def test_lock_user_by_numeric_id_string(db_session_and_client):
    client, engine = db_session_and_client
    payload = {"user_id": "2", "reason": "Khóa theo ID số"}

    response = client.post("/api/users/lock", json=payload)

    assert response.status_code == 200
    assert response.json()["success"] is True

    # Xác nhận is_active của user 2 trong DB thành False
    with engine.connect() as conn:
        status_after = conn.execute(text("SELECT is_active FROM users WHERE id = 2")).scalar()
        assert bool(status_after) is False


# =====================================================================
# 3. Test mở khóa tài khoản thành công và xác nhận UPDATE is_active = True
# =====================================================================
def test_unlock_user_success_updates_is_active_to_true(db_session_and_client):
    client, engine = db_session_and_client

    # Khóa user trước
    client.post("/api/users/lock", json={"user_id": "sales01", "reason": "Tạm khóa"})
    with engine.connect() as conn:
        assert bool(conn.execute(text("SELECT is_active FROM users WHERE username = 'sales01'")).scalar()) is False

    # Gọi API mở khóa tài khoản bằng username
    response = client.post("/api/users/unlock/sales01")

    # Xác nhận response thành công và không bị 500
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "đã được mở khóa thành công" in data["message"]

    # Xác nhận trực tiếp trong Database: is_active đã được cập nhật thành True (1)
    with engine.connect() as conn:
        status_after = conn.execute(text("SELECT is_active FROM users WHERE username = 'sales01'")).scalar()
        assert bool(status_after) is True, f"Expected is_active=True, got {status_after}"


def test_unlock_user_by_numeric_id(db_session_and_client):
    client, engine = db_session_and_client

    # Khóa user 2 trước
    client.post("/api/users/lock", json={"user_id": "2", "reason": "Tạm khóa"})

    # Mở khóa bằng ID số "2"
    response = client.post("/api/users/unlock/2")

    assert response.status_code == 200
    assert response.json()["success"] is True

    with engine.connect() as conn:
        status_after = conn.execute(text("SELECT is_active FROM users WHERE id = 2")).scalar()
        assert bool(status_after) is True


# =====================================================================
# 4. Test khi tài khoản không tồn tại trong Database (Lỗi 404)
# =====================================================================
def test_lock_nonexistent_user_returns_404(db_session_and_client):
    client, _ = db_session_and_client
    payload = {"user_id": "nonexistent_user", "reason": "Lý do khóa"}

    response = client.post("/api/users/lock", json=payload)

    assert response.status_code == 404
    assert "Không tìm thấy người dùng" in response.json()["detail"]


def test_unlock_nonexistent_user_returns_404(db_session_and_client):
    client, _ = db_session_and_client

    response = client.post("/api/users/unlock/nonexistent_user")

    assert response.status_code == 404
    assert "Không tìm thấy người dùng" in response.json()["detail"]


def test_bulk_import_users_from_excel_creates_records_and_skips_duplicates():
    from src.backend import models
    from src.backend.users_service import bulk_create_users_from_rows

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"
    sheet.append(["full_name", "username", "email", "phone", "role_codes"])
    sheet.append(["Nguyễn Văn A", "nguyenvana", "a@example.com", "0901234567", "SALES_REP"])
    sheet.append(["Trần Thị B", "tranthib", "b@example.com", "0901234568", "WAREHOUSE"])
    sheet.append(["Nguyễn Văn A", "nguyenvana", "a2@example.com", "0901234569", "SALES_REP"])

    rows = []
    for row in sheet.iter_rows(values_only=True):
        if row[0] is None:
            continue
        rows.append(dict(zip(["full_name", "username", "email", "phone", "role_codes"], row)))

    db = SimpleNamespace()
    db.query = lambda *args, **kwargs: None

    # Use a real in-memory database to validate actual persistence.
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    engine = create_engine("sqlite:///:memory:")
    models.Base.metadata.create_all(bind=engine)
    SessionLocal = sessionmaker(bind=engine)
    session = SessionLocal()

    try:
        created, errors, duplicates = bulk_create_users_from_rows(session, rows, actor_user_id=99)
        assert created == 2
        assert duplicates == 1
        assert len(errors) == 0
        assert session.query(models.User).count() == 2
        assert {user.username for user in session.query(models.User).all()} == {"nguyenvana", "tranthib"}
    finally:
        session.close()


def test_admin_user_import_preview_reports_rows_without_creating_users():
    from src.backend import models
    from src.backend.database import get_db
    from src.backend.users import router as users_router

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    models.Base.metadata.create_all(bind=engine)
    TestingSession = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    app = FastAPI()
    app.include_router(users_router)

    def override_get_db():
        session = TestingSession()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[require_admin] = lambda: SimpleNamespace(id=999, roles=[SimpleNamespace(code="ADMIN")])
    content = (
        "full_name,username,email,phone,role_codes\n"
        "Valid User,valid_user,valid@example.com,0901234567,SALES_REP\n"
        "Invalid User,invalid_user,not-an-email,0901234568,SALES_REP\n"
        "Duplicate User,valid_user,duplicate@example.com,0901234569,SALES_REP\n"
    )

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/admin/users/import/preview",
            files={"file": ("users.csv", content, "text/csv")},
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] == 3
    assert payload["valid"] == 1
    assert payload["errors"] == 2
    assert payload["duplicates"] == 1
    assert [row["valid"] for row in payload["rows"]] == [True, False, False]
    assert "Email không hợp lệ" in payload["rows"][1]["error"]
    with TestingSession() as session:
        assert session.query(models.User).count() == 0


def test_admin_user_import_template_downloads_valid_xlsx():
    from src.backend.users import router as users_router

    app = FastAPI()
    app.include_router(users_router)
    app.dependency_overrides[require_admin] = lambda: SimpleNamespace(id=999, roles=[SimpleNamespace(code="ADMIN")])

    with TestClient(app) as client:
        response = client.get("/api/v1/admin/users/import/template")

    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    assert list(workbook.active.iter_rows(values_only=True))[0] == ("Họ tên", "Tên đăng nhập", "Email", "Số điện thoại", "Vai trò")


def test_admin_user_import_skips_invalid_rows_and_creates_valid_rows(monkeypatch):
    from src.backend import models
    from src.backend.database import get_db
    from src.backend.users import router as users_router
    from src.backend.email_service import email_service

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    models.Base.metadata.create_all(bind=engine)
    TestingSession = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    app = FastAPI()
    app.include_router(users_router)

    def override_get_db():
        session = TestingSession()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[require_admin] = lambda: SimpleNamespace(id=999, roles=[SimpleNamespace(code="ADMIN")])
    monkeypatch.setattr(email_service, "send_activation_email", lambda **kwargs: True)
    content = (
        "full_name,username,email,phone,role_codes\n"
        "Valid User,valid_user,valid@example.com,0901234567,SALES_REP\n"
        "Invalid User,invalid_user,invalid-email,0901234568,SALES_REP\n"
    )

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/admin/users/import",
            files={"file": ("users.csv", content, "text/csv")},
        )

    assert response.status_code == 200
    result = response.json()
    assert result["created"] == 1
    assert result["skipped"] == 1
    assert "Dòng 3" in result["errors"][0]
    with TestingSession() as session:
        assert session.query(models.User).count() == 1
        assert session.query(models.User).filter_by(username="valid_user").one()
