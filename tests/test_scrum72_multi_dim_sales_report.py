"""Unit & Integration Tests for SCRUM-72 (Epic SCRUM-91): Multi-dimensional Sales Report.

Verifies:
1. Multi-dimensional grouping (staff/sales_rep, region/territory, customer/dealer, category, product/sku).
2. Accuracy of metrics: Revenue, Volume (Quantity), COGS, and Gross Margin ((Revenue - COGS) / Revenue * 100).
3. Data RBAC confidentiality:
   - Sales Manager and Admin see COGS, gross profit, and gross margin.
   - Sales Rep and other roles have COGS and gross margin masked/hidden (null/None).
4. Period-over-Period (PoP) comparison and growth rates calculation (% revenue growth, % quantity growth).
5. Excel export (.xlsx) generation with proper MIME type and data structure.
6. Python 3.8 compatibility and zero regression across test suite.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from typing import Any, Dict, Generator, Tuple

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.v1.endpoints.rbac import create_access_token
from app.core.database import get_db
from app.models.models import (
    AccountStatus,
    Base,
    Category,
    Customer,
    Order,
    OrderItem,
    OrderStatus,
    Product,
    Role,
    RoleCode,
    Territory,
    User,
)
from main import app


def make_auth_header(role: str, username: str = "test_user") -> Dict[str, str]:
    """Helper to generate JWT Bearer Authorization header."""
    token = create_access_token({"sub": username, "role": role})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def scrum72_env() -> Generator[Tuple[TestClient, sessionmaker], None, None]:
    """Setup isolated in-memory SQLite database with rich sales report test data."""
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    def override_get_db() -> Generator[Session, None, None]:
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db

    # Seed data
    with TestingSessionLocal() as db:
        # 1. Roles & Users
        manager_role = Role(code=RoleCode.SALES_MANAGER.value, name="Quản lý kinh doanh")
        sales_role = Role(code=RoleCode.SALES_REP.value, name="Nhân viên kinh doanh")
        admin_role = Role(code=RoleCode.ADMIN.value, name="Quản trị hệ thống")
        db.add_all([manager_role, sales_role, admin_role])
        db.flush()

        user_mgr = User(
            username="manager_01",
            email="manager@company.com",
            full_name="Quản Lý Nguyễn",
            password_hash="fakehash",
            status=AccountStatus.ACTIVE,
            roles=[manager_role],
        )
        user_rep_a = User(
            username="rep_a",
            email="repa@company.com",
            full_name="Trần Văn A",
            password_hash="fakehash",
            status=AccountStatus.ACTIVE,
            roles=[sales_role],
        )
        user_rep_b = User(
            username="rep_b",
            email="repb@company.com",
            full_name="Lê Thị B",
            password_hash="fakehash",
            status=AccountStatus.ACTIVE,
            roles=[sales_role],
        )
        user_admin = User(
            username="admin_01",
            email="admin@company.com",
            full_name="Admin Hệ Thống",
            password_hash="fakehash",
            status=AccountStatus.ACTIVE,
            roles=[admin_role],
        )
        db.add_all([user_mgr, user_rep_a, user_rep_b, user_admin])
        db.flush()

        # 2. Territories
        t_bac = Territory(code="MB", name="Miền Bắc")
        t_nam = Territory(code="MN", name="Miền Nam")
        db.add_all([t_bac, t_nam])
        db.flush()

        # 3. Categories
        cat_bev = Category(code="BEV", name="Đồ uống", description="Nước ngọt, bia")
        cat_food = Category(code="FOOD", name="Thực phẩm", description="Bánh kẹo, đồ ăn")
        db.add_all([cat_bev, cat_food])
        db.flush()

        # 4. Products with Cost Prices
        # Product 1: BEV-COCA (Sale: 10,000, Cost: 6,000 -> Gross Margin = 40%)
        p_coca = Product(
            sku="BEV-COCA",
            name="Coca Cola 330ml",
            category_id=cat_bev.id,
            base_unit="Lon",
            base_price=10000.0,
            cost_price=6000.0,
            status="ACTIVE",
        )
        # Product 2: BEV-PEPSI (Sale: 10,000, Cost: 7,000 -> Gross Margin = 30%)
        p_pepsi = Product(
            sku="BEV-PEPSI",
            name="Pepsi 330ml",
            category_id=cat_bev.id,
            base_unit="Lon",
            base_price=10000.0,
            cost_price=7000.0,
            status="ACTIVE",
        )
        # Product 3: FOOD-SNACK (Sale: 20,000, Cost: 12,000 -> Gross Margin = 40%)
        p_snack = Product(
            sku="FOOD-SNACK",
            name="Snack Khoai Tây",
            category_id=cat_food.id,
            base_unit="Gói",
            base_price=20000.0,
            cost_price=12000.0,
            status="ACTIVE",
        )
        db.add_all([p_coca, p_pepsi, p_snack])
        db.flush()

        # 5. Customers
        # Customer 1: Miền Bắc, Rep A
        cust1 = Customer(
            code="DL-HN-01",
            name="Đại lý Cầu Giấy",
            region="Miền Bắc",
            territory_id=t_bac.id,
            sales_rep_id=user_rep_a.id,
        )
        # Customer 2: Miền Nam, Rep B
        cust2 = Customer(
            code="DL-SG-01",
            name="Đại lý Quận 1",
            region="Miền Nam",
            territory_id=t_nam.id,
            sales_rep_id=user_rep_b.id,
        )
        db.add_all([cust1, cust2])
        db.flush()

        # 6. Orders: Current Period (October 2026)
        oct_date = datetime(2026, 10, 15, 10, 0, 0)
        # Order 1: Cust 1 (Rep A, MB). Coca: 100 lon @ 10,000 = 1,000,000 (Cost: 600,000). Snack: 50 gói @ 20,000 = 1,000,000 (Cost: 600,000). Total: 2,000,000 (Cost: 1,200,000, Margin: 40%)
        ord1 = Order(
            order_code="ORD-OCT-01",
            customer_id=cust1.id,
            created_by_id=user_rep_a.id,
            status=OrderStatus.COMPLETED,
            total_amount=2000000.0,
            created_at=oct_date,
        )
        db.add(ord1)
        db.flush()
        item1_1 = OrderItem(
            order_id=ord1.id,
            product_id=p_coca.id,
            sku="BEV-COCA",
            product_name="Coca Cola 330ml",
            quantity=100,
            unit_price=10000.0,
            line_total=1000000.0,
        )
        item1_2 = OrderItem(
            order_id=ord1.id,
            product_id=p_snack.id,
            sku="FOOD-SNACK",
            product_name="Snack Khoai Tây",
            quantity=50,
            unit_price=20000.0,
            line_total=1000000.0,
        )
        db.add_all([item1_1, item1_2])

        # Order 2: Cust 2 (Rep B, MN). Pepsi: 200 lon @ 10,000 = 2,000,000 (Cost: 1,400,000, Profit: 600,000, Margin: 30%)
        ord2 = Order(
            order_code="ORD-OCT-02",
            customer_id=cust2.id,
            created_by_id=user_rep_b.id,
            status=OrderStatus.APPROVED,
            total_amount=2000000.0,
            created_at=oct_date,
        )
        db.add(ord2)
        db.flush()
        item2_1 = OrderItem(
            order_id=ord2.id,
            product_id=p_pepsi.id,
            sku="BEV-PEPSI",
            product_name="Pepsi 330ml",
            quantity=200,
            unit_price=10000.0,
            line_total=2000000.0,
        )
        db.add(item2_1)

        # Order 3: CANCELLED order - MUST NOT be counted in revenue/quantity
        ord_cancelled = Order(
            order_code="ORD-OCT-CANCELLED",
            customer_id=cust1.id,
            created_by_id=user_rep_a.id,
            status=OrderStatus.CANCELLED,
            total_amount=5000000.0,
            created_at=oct_date,
        )
        db.add(ord_cancelled)
        db.flush()
        item_canc = OrderItem(
            order_id=ord_cancelled.id,
            product_id=p_coca.id,
            sku="BEV-COCA",
            quantity=500,
            unit_price=10000.0,
            line_total=5000000.0,
        )
        db.add(item_canc)

        # 7. Orders: Previous Period (September 2026) for Comparison
        sep_date = datetime(2026, 9, 15, 10, 0, 0)
        # Order 4: Cust 1 (Rep A, MB). Coca: 80 lon @ 10,000 = 800,000. Snack: 40 gói @ 20,000 = 800,000. Total: 1,600,000
        ord4 = Order(
            order_code="ORD-SEP-01",
            customer_id=cust1.id,
            created_by_id=user_rep_a.id,
            status=OrderStatus.COMPLETED,
            total_amount=1600000.0,
            created_at=sep_date,
        )
        db.add(ord4)
        db.flush()
        item4_1 = OrderItem(
            order_id=ord4.id,
            product_id=p_coca.id,
            sku="BEV-COCA",
            quantity=80,
            unit_price=10000.0,
            line_total=800000.0,
        )
        item4_2 = OrderItem(
            order_id=ord4.id,
            product_id=p_snack.id,
            sku="FOOD-SNACK",
            quantity=40,
            unit_price=20000.0,
            line_total=800000.0,
        )
        db.add_all([item4_1, item4_2])

        # Order 5: Cust 2 (Rep B, MN). Pepsi: 150 lon @ 10,000 = 1,500,000
        ord5 = Order(
            order_code="ORD-SEP-02",
            customer_id=cust2.id,
            created_by_id=user_rep_b.id,
            status=OrderStatus.COMPLETED,
            total_amount=1500000.0,
            created_at=sep_date,
        )
        db.add(ord5)
        db.flush()
        item5_1 = OrderItem(
            order_id=ord5.id,
            product_id=p_pepsi.id,
            sku="BEV-PEPSI",
            quantity=150,
            unit_price=10000.0,
            line_total=1500000.0,
        )
        db.add(item5_1)

        db.commit()

    with TestClient(app) as client:
        yield client, TestingSessionLocal

    app.dependency_overrides.clear()
    Base.metadata.drop_all(bind=engine)


# ==============================================================================
# 1. KIỂM THỬ GOM NHÓM ĐA CHIỀU (MULTI-DIMENSIONAL GROUP BY)
# ==============================================================================

def test_sales_analytics_group_by_product_sku(scrum72_env):
    """Gom nhóm theo Mặt hàng cụ thể (SKU): Tính đúng doanh số, sản lượng và biên lợi nhuận."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager")

    res = client.get(
        "/api/v1/reports/sales-analytics?group_by=sku&from_date=2026-10-01&to_date=2026-10-31",
        headers=headers,
    )
    assert res.status_code == 200
    body = res.json()
    assert body["success"] is True
    data = body["data"]
    assert data["group_by"] == "product"

    summary = data["summary"]
    # Total valid October revenue: 2,000,000 (ord1) + 2,000,000 (ord2) = 4,000,000
    # Total quantity: 100 (coca) + 50 (snack) + 200 (pepsi) = 350
    # Total COGS: (100 * 6000) + (50 * 12000) + (200 * 7000) = 600,000 + 600,000 + 1,400,000 = 2,600,000
    # Gross Margin: (4,000,000 - 2,600,000) / 4,000,000 * 100 = 35.0%
    assert summary["total_revenue"] == 4000000.0
    assert summary["total_quantity"] == 350
    assert summary["total_cogs"] == 2600000.0
    assert summary["gross_profit"] == 1400000.0
    assert summary["gross_margin"] == 35.0

    items_map = {item["code"]: item for item in data["items"]}
    # Coca: 100 units, Rev 1,000,000, COGS 600,000, Profit 400,000, Margin 40%
    coca = items_map["BEV-COCA"]
    assert coca["quantity"] == 100
    assert coca["revenue"] == 1000000.0
    assert coca["cogs"] == 600000.0
    assert coca["gross_profit"] == 400000.0
    assert coca["gross_margin"] == 40.0

    # Pepsi: 200 units, Rev 2,000,000, COGS 1,400,000, Profit 600,000, Margin 30%
    pepsi = items_map["BEV-PEPSI"]
    assert pepsi["quantity"] == 200
    assert pepsi["revenue"] == 2000000.0
    assert pepsi["cogs"] == 1400000.0
    assert pepsi["gross_profit"] == 600000.0
    assert pepsi["gross_margin"] == 30.0

    # Snack: 50 units, Rev 1,000,000, COGS 600,000, Profit 400,000, Margin 40%
    snack = items_map["FOOD-SNACK"]
    assert snack["quantity"] == 50
    assert snack["revenue"] == 1000000.0
    assert snack["cogs"] == 600000.0
    assert snack["gross_profit"] == 400000.0
    assert snack["gross_margin"] == 40.0


def test_sales_analytics_group_by_category(scrum72_env):
    """Gom nhóm theo Nhóm hàng (Category): BEV vs FOOD."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager")

    res = client.get(
        "/api/v1/reports/sales-analytics?group_by=category&from_date=2026-10-01&to_date=2026-10-31",
        headers=headers,
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["group_by"] == "category"

    items_map = {item["code"]: item for item in data["items"]}
    # BEV: Coca (1M) + Pepsi (2M) = 3M, Qty 300, COGS 600k + 1.4M = 2.0M, Margin = 33.33%
    bev = items_map["BEV"]
    assert bev["revenue"] == 3000000.0
    assert bev["quantity"] == 300
    assert bev["cogs"] == 2000000.0
    assert bev["gross_profit"] == 1000000.0
    assert round(bev["gross_margin"], 2) == 33.33

    # FOOD: Snack 1M, Qty 50, COGS 600k, Margin = 40%
    food = items_map["FOOD"]
    assert food["revenue"] == 1000000.0
    assert food["quantity"] == 50
    assert food["cogs"] == 600000.0
    assert food["gross_margin"] == 40.0


def test_sales_analytics_group_by_region(scrum72_env):
    """Gom nhóm theo Khu vực (Region / Territory): Miền Bắc vs Miền Nam."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager")

    res = client.get(
        "/api/v1/reports/sales-analytics?group_by=region&from_date=2026-10-01&to_date=2026-10-31",
        headers=headers,
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["group_by"] == "region"

    items_map = {item["name"]: item for item in data["items"]}
    # Miền Bắc (Cust 1): Rev 2M, Qty 150, COGS 1.2M, Margin 40%
    mb = items_map["Miền Bắc"]
    assert mb["revenue"] == 2000000.0
    assert mb["quantity"] == 150
    assert mb["cogs"] == 1200000.0
    assert mb["gross_margin"] == 40.0

    # Miền Nam (Cust 2): Rev 2M, Qty 200, COGS 1.4M, Margin 30%
    mn = items_map["Miền Nam"]
    assert mn["revenue"] == 2000000.0
    assert mn["quantity"] == 200
    assert mn["cogs"] == 1400000.0
    assert mn["gross_margin"] == 30.0


def test_sales_analytics_group_by_staff(scrum72_env):
    """Gom nhóm theo Nhân viên kinh doanh (Staff / Sales Rep): Rep A vs Rep B."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager")

    res = client.get(
        "/api/v1/reports/sales-analytics?group_by=staff&from_date=2026-10-01&to_date=2026-10-31",
        headers=headers,
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["group_by"] == "staff"

    items_map = {item["code"]: item for item in data["items"]}
    rep_a = items_map["rep_a"]
    assert rep_a["revenue"] == 2000000.0
    assert rep_a["quantity"] == 150
    assert rep_a["cogs"] == 1200000.0
    assert rep_a["gross_margin"] == 40.0

    rep_b = items_map["rep_b"]
    assert rep_b["revenue"] == 2000000.0
    assert rep_b["quantity"] == 200
    assert rep_b["cogs"] == 1400000.0
    assert rep_b["gross_margin"] == 30.0


def test_sales_analytics_group_by_customer(scrum72_env):
    """Gom nhóm theo Đại lý (Customer / Dealer)."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager")

    res = client.get(
        "/api/v1/reports/sales-analytics?group_by=customer&from_date=2026-10-01&to_date=2026-10-31",
        headers=headers,
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["group_by"] == "customer"

    items_map = {item["code"]: item for item in data["items"]}
    cust1 = items_map["DL-HN-01"]
    assert cust1["name"] == "Đại lý Cầu Giấy"
    assert cust1["revenue"] == 2000000.0
    assert cust1["quantity"] == 150

    cust2 = items_map["DL-SG-01"]
    assert cust2["name"] == "Đại lý Quận 1"
    assert cust2["revenue"] == 2000000.0
    assert cust2["quantity"] == 200


# ==============================================================================
# 2. KIỂM THỬ PHÂN QUYỀN RBAC (DATA CONFIDENTIALITY)
# ==============================================================================

def test_sales_manager_can_see_gross_margin_and_cogs(scrum72_env):
    """Quản lý kinh doanh xem đầy đủ biên lợi nhuận, giá vốn và lợi nhuận gộp."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager", username="manager_01")

    res = client.get("/api/v1/reports/sales-analytics?group_by=category", headers=headers)
    assert res.status_code == 200
    data = res.json()["data"]

    # Manager must see all financial figures
    assert data["summary"]["total_cogs"] is not None
    assert data["summary"]["gross_profit"] is not None
    assert data["summary"]["gross_margin"] is not None

    for item in data["items"]:
        assert item["cogs"] is not None
        assert item["gross_profit"] is not None
        assert item["gross_margin"] is not None


def test_admin_can_see_gross_margin_and_cogs(scrum72_env):
    """Admin xem được biên lợi nhuận và giá vốn."""
    client, _ = scrum72_env
    headers = make_auth_header("Admin", username="admin_01")

    res = client.get("/api/v1/reports/sales-analytics?group_by=category", headers=headers)
    assert res.status_code == 200
    data = res.json()["data"]

    assert data["summary"]["total_cogs"] is not None
    assert data["summary"]["gross_profit"] is not None
    assert data["summary"]["gross_margin"] is not None


def test_sales_rep_has_cogs_and_gross_margin_masked(scrum72_env):
    """Nhân viên kinh doanh thông thường (Sales Rep) bị ẩn/chặn chỉ số biên lợi nhuận và giá vốn (trả về null)."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Rep", username="rep_a")

    res = client.get("/api/v1/reports/sales-analytics?group_by=category", headers=headers)
    assert res.status_code == 200
    data = res.json()["data"]

    # Revenue and quantity are visible
    assert data["summary"]["total_revenue"] > 0
    assert data["summary"]["total_quantity"] > 0

    # Financial margin data MUST be null/None for regular staff
    assert data["summary"]["total_cogs"] is None
    assert data["summary"]["gross_profit"] is None
    assert data["summary"]["gross_margin"] is None

    for item in data["items"]:
        assert item["revenue"] > 0
        assert item["quantity"] > 0
        assert item["cogs"] is None
        assert item["gross_profit"] is None
        assert item["gross_margin"] is None


def test_unauthenticated_request_rejected(scrum72_env):
    """Yêu cầu không có token bị từ chối 401 Unauthorized."""
    client, _ = scrum72_env
    res = client.get("/api/v1/reports/sales-analytics")
    assert res.status_code == 401


# ==============================================================================
# 3. KIỂM THỬ SO SÁNH CÙNG KỲ TRƯỚC (PERIOD-OVER-PERIOD)
# ==============================================================================

def test_period_over_period_comparison_growth_calculation(scrum72_env):
    """So sánh tháng 10/2026 với tháng 9/2026: Doanh số & sản lượng cùng kỳ, % tăng trưởng."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager")

    res = client.get(
        "/api/v1/reports/sales-analytics"
        "?group_by=category"
        "&from_date=2026-10-01&to_date=2026-10-31"
        "&compare_to=previous_period",
        headers=headers,
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["comparison"] is not None

    summary = data["summary"]
    # October Total Rev: 4,000,000, Qty: 350
    # September Total Rev: 1,600,000 (ord4) + 1,500,000 (ord5) = 3,100,000. Qty: 80 + 40 + 150 = 270
    assert summary["total_revenue"] == 4000000.0
    assert summary["previous_revenue"] == 3100000.0
    # Growth = ((4,000,000 - 3,100,000) / 3,100,000) * 100 = 29.03%
    expected_rev_growth = round(((4000000.0 - 3100000.0) / 3100000.0) * 100.0, 2)
    assert summary["revenue_growth_rate"] == expected_rev_growth

    # Quantity Growth = ((350 - 270) / 270) * 100 = 29.63%
    expected_qty_growth = round(((350 - 270) / 270.0) * 100.0, 2)
    assert summary["quantity_growth_rate"] == expected_qty_growth

    # Detailed item growth: Category BEV
    items_map = {item["code"]: item for item in data["items"]}
    bev = items_map["BEV"]
    # October BEV: Rev 3,000,000, Qty 300
    # September BEV: Coca 800,000 (80 lon) + Pepsi 1,500,000 (150 lon) = 2,300,000 (230 lon)
    assert bev["revenue"] == 3000000.0
    assert bev["previous_revenue"] == 2300000.0
    expected_bev_growth = round(((3000000.0 - 2300000.0) / 2300000.0) * 100.0, 2)
    assert bev["revenue_growth_rate"] == expected_bev_growth


# ==============================================================================
# 4. KIỂM THỬ XUẤT FILE EXCEL (.XLSX)
# ==============================================================================

def test_export_sales_analytics_excel_for_manager(scrum72_env):
    """Xuất báo cáo Excel cho Quản lý kinh doanh: MIME type chuẩn, đầy đủ cột Doanh số, Sản lượng, Giá vốn, Biên lợi nhuận."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager")

    res = client.get(
        "/api/v1/reports/sales-analytics/export?group_by=category&from_date=2026-10-01&to_date=2026-10-31&compare_to=previous_period",
        headers=headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "sales_analytics_category_" in res.headers.get("content-disposition", "")

    # Read binary content into openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    assert ws.title == "Doanh Số Đa Chiều"

    # Check headers at row 5
    header_values = [cell.value for cell in ws[5] if cell.value is not None]
    assert "STT" in header_values
    assert "Mã" in header_values
    assert "Doanh số (VNĐ)" in header_values
    assert "Sản lượng bán" in header_values
    # Manager must have Margin and COGS columns
    assert "Giá vốn COGS (VNĐ)" in header_values
    assert "Lợi nhuận gộp (VNĐ)" in header_values
    assert "Biên lợi nhuận (%)" in header_values
    assert "Doanh số kỳ trước (VNĐ)" in header_values
    assert "Tăng trưởng DS (%)" in header_values

    # Check total row exists at bottom
    last_row_vals = [cell.value for cell in ws[ws.max_row] if cell.value is not None]
    assert "TỔNG CỘNG" in last_row_vals
    assert 4000000.0 in last_row_vals  # Total revenue in total row


def test_export_sales_analytics_excel_for_sales_rep_masks_financials(scrum72_env):
    """Xuất báo cáo Excel cho Nhân viên kinh doanh: Cột Giá vốn COGS và Biên lợi nhuận bị loại trừ hoàn toàn."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Rep")

    res = client.get(
        "/api/v1/reports/sales-analytics/export?group_by=sku&from_date=2026-10-01&to_date=2026-10-31",
        headers=headers,
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    header_values = [cell.value for cell in ws[5] if cell.value is not None]

    assert "Doanh số (VNĐ)" in header_values
    assert "Sản lượng bán" in header_values
    # Confidential columns MUST NOT appear in Sales Rep export
    assert "Giá vốn COGS (VNĐ)" not in header_values
    assert "Lợi nhuận gộp (VNĐ)" not in header_values
    assert "Biên lợi nhuận (%)" not in header_values


def test_sales_report_compat_endpoints(scrum72_env):
    """Kiểm tra tính tương thích: /reports/sales-analytics và /api/v1/reports/sales."""
    client, _ = scrum72_env
    headers = make_auth_header("Sales Manager")

    # 1. Compat route without /api/v1
    res1 = client.get("/reports/sales-analytics?group_by=category", headers=headers)
    assert res1.status_code == 200
    assert res1.json()["data"]["group_by"] == "category"

    # 2. Alias route /api/v1/reports/sales
    res2 = client.get("/api/v1/reports/sales?group_by=region", headers=headers)
    assert res2.status_code == 200
    assert res2.json()["data"]["group_by"] == "region"
