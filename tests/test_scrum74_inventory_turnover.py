"""Tests for SCRUM-74: Quản lý kho xem báo cáo tồn kho và vòng quay hàng để phát hiện hàng chậm luân chuyển.

Yêu cầu nghiệp vụ (Jira SCRUM-74):
1. Giá trị tồn theo kho và theo nhóm hàng (valuation by warehouse and category).
2. Số ngày tồn bình quân của từng SKU (DSI) và tỷ lệ vòng quay hàng (Turnover).
3. Phát hiện và chỉ ra các mặt hàng chậm luân chuyển đang chôn vốn (is_slow_moving = True).
4. Bộ lọc theo kho, theo nhóm hàng, khoảng thời gian kỳ báo cáo, và tùy chọn chỉ xem hàng chậm luân chuyển.
5. Xuất báo cáo ra file Excel (.xlsx) chuẩn với 2 sheets định dạng chuyên nghiệp.
6. Phân quyền: Quản lý kho (WH Manager), Quản trị (Admin), Nhân viên kho, Quản lý kinh doanh được phép xem và xuất báo cáo.
"""
import io
import pytest
import openpyxl
from datetime import datetime, timezone, timedelta
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import get_db
from app.models.models import (
    Base,
    Category,
    Customer,
    DeliveryAddress,
    Dispatch,
    DispatchItem,
    Inventory,
    Order,
    OrderItem,
    OrderStatus,
    Product,
    RoleCode,
    User,
    Warehouse,
)
from app.api.v1.endpoints.rbac import create_access_token
from app.api.v1.endpoints.reports import calculate_inventory_turnover_data
from backend.main import app


def make_auth_header(role: str = "Warehouse Manager", username: str = "wh_mgr") -> dict[str, str]:
    token = create_access_token({"sub": username, "role": role})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def scrum74_env():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db

    with TestingSessionLocal() as db:
        # Tạo 2 kho hàng
        wh1 = Warehouse(id=1, code="WH-HN", name="Kho Tổng Hà Nội", is_active=True)
        wh2 = Warehouse(id=2, code="WH-HCM", name="Kho Chi Nhánh TP.HCM", is_active=True)
        db.add_all([wh1, wh2])

        # Tạo 2 nhóm hàng
        cat1 = Category(id=1, code="BEV", name="Đồ uống & Nước giải khát", description="Nước ngọt, bia, trà")
        cat2 = Category(id=2, code="FOOD", name="Thực phẩm khô", description="Bánh kẹo, mì gói, gia vị")
        db.add_all([cat1, cat2])

        # Tạo 3 sản phẩm:
        # P1: Bán chạy (Fast moving)
        p1 = Product(id=1, sku="SKU-COCA", name="Coca Cola Lon 330ml", category_id=1, base_price=10000.0, cost_price=8000.0, base_unit="Lon", status="ACTIVE")
        # P2: Chậm luân chuyển (Slow moving: tồn nhiều, bán rất ít)
        p2 = Product(id=2, sku="SKU-RUOU", name="Rượu Vang Thượng Hạng", category_id=1, base_price=500000.0, cost_price=400000.0, base_unit="Chai", status="ACTIVE")
        # P3: Tồn đọng không có giao dịch bán nào (Dead stock)
        p3 = Product(id=3, sku="SKU-BAP", name="Bắp Rang Bơ Nhập Khẩu", category_id=2, base_price=60000.0, cost_price=45000.0, base_unit="Gói", status="ACTIVE")
        db.add_all([p1, p2, p3])

        # Tồn kho tại Kho HN (WH 1):
        # P1: 100 lon @ 8k = 800k
        # P2: 50 chai @ 400k = 20tr
        # P3: 200 gói @ 45k = 9tr
        inv1 = Inventory(product_id=1, warehouse_id=1, physical_qty=100.0, reserved_qty=10.0, available_qty=90.0)
        inv2 = Inventory(product_id=2, warehouse_id=1, physical_qty=50.0, reserved_qty=0.0, available_qty=50.0)
        inv3 = Inventory(product_id=3, warehouse_id=1, physical_qty=200.0, reserved_qty=0.0, available_qty=200.0)

        # Tồn kho tại Kho HCM (WH 2):
        # P1: 300 lon @ 8k = 2.4tr
        inv4 = Inventory(product_id=1, warehouse_id=2, physical_qty=300.0, reserved_qty=50.0, available_qty=250.0)
        db.add_all([inv1, inv2, inv3, inv4])

        # Giả lập lịch sử xuất kho trong kỳ 90 ngày:
        # P1 tại WH1: xuất 200 lon (vòng quay tốt)
        # P2 tại WH1: xuất 2 chai (vòng quay cực chậm: 50 / (2/90) = 2250 ngày)
        # P3 tại WH1: xuất 0 gói (chậm luân chuyển / dead stock)
        now = datetime.now(timezone.utc)
        disp1 = Dispatch(
            dispatch_code="DSP-PAST-01",
            order_id=1,
            warehouse_id=1,
            is_confirmed=True,
            confirmed_at=now - timedelta(days=20),
        )
        db.add(disp1)
        db.flush()
        db.add(DispatchItem(dispatch_id=disp1.id, product_id=1, sku="SKU-COCA", requested_qty=200, actual_qty=200, missing_qty=0))
        db.add(DispatchItem(dispatch_id=disp1.id, product_id=2, sku="SKU-RUOU", requested_qty=2, actual_qty=2, missing_qty=0))

        # P1 tại WH2: xuất 600 lon (vòng quay tốt)
        disp2 = Dispatch(
            dispatch_code="DSP-PAST-02",
            order_id=2,
            warehouse_id=2,
            is_confirmed=True,
            confirmed_at=now - timedelta(days=10),
        )
        db.add(disp2)
        db.flush()
        db.add(DispatchItem(dispatch_id=disp2.id, product_id=1, sku="SKU-COCA", requested_qty=600, actual_qty=600, missing_qty=0))

        # User
        mgr = User(id=601, username="wh_mgr", email="whmgr@oms.com", full_name="Quản Lý Kho Trưởng", password_hash="h", is_active=True)
        db.add(mgr)
        db.commit()

    client = TestClient(app)
    yield client, TestingSessionLocal

    app.dependency_overrides.clear()
    Base.metadata.drop_all(bind=engine)


def test_inventory_valuation_by_warehouse_and_category(scrum74_env):
    """AC1: Tính chính xác giá trị tồn theo kho và theo nhóm hàng."""
    client, SessionLocal = scrum74_env
    headers = make_auth_header("Warehouse Manager", "wh_mgr")

    res = client.get("/api/v1/reports/inventory-turnover", headers=headers)
    assert res.status_code == 200
    data = res.json()["data"]

    # Tổng giá trị tồn kho toàn hệ thống:
    # WH1: P1(100*8k=800k) + P2(50*400k=20tr) + P3(200*45k=9tr) = 29.800.000 VND
    # WH2: P1(300*8k=2.4tr)
    # Tổng cộng = 32.200.000 VND
    assert data["total_inventory_value"] == 32200000.0

    # Phân bổ theo kho:
    by_wh = {w["warehouse_code"]: w for w in data["by_warehouse"]}
    assert by_wh["WH-HN"]["total_value"] == 29800000.0
    assert by_wh["WH-HN"]["total_items"] == 3
    assert by_wh["WH-HCM"]["total_value"] == 2400000.0
    assert by_wh["WH-HCM"]["total_items"] == 1

    # Phân bổ theo nhóm hàng:
    by_cat = {c["category_name"]: c for c in data["by_category"]}
    # Đồ uống: P1 ở cả 2 kho (800k + 2.4tr = 3.2tr) + P2 (20tr) = 23.200.000 VND
    # Thực phẩm khô: P3 (9.000.000 VND)
    assert by_cat["Đồ uống & Nước giải khát"]["total_value"] == 23200000.0
    assert by_cat["Thực phẩm khô"]["total_value"] == 9000000.0


def test_inventory_turnover_and_slow_moving_identification(scrum74_env):
    """AC2 & AC3: Số ngày tồn bình quân của từng SKU và chỉ ra hàng chậm luân chuyển."""
    client, SessionLocal = scrum74_env
    headers = make_auth_header("Warehouse Manager", "wh_mgr")

    res = client.get("/api/v1/reports/inventory-turnover?period_days=90&threshold_days=90", headers=headers)
    assert res.status_code == 200
    data = res.json()["data"]

    sku_items = {f"{it['sku']}_{it['warehouse_code']}": it for it in data["sku_details"]}

    # P1 tại WH1: xuất 200 trong 90 ngày (daily_rate = 200/90 = 2.22). Tồn 100 -> DSI = 100 / 2.22 = 45 ngày (< 90) -> Không chậm
    coca_hn = sku_items["SKU-COCA_WH-HN"]
    assert coca_hn["is_slow_moving"] is False
    assert coca_hn["turnover_ratio"] == 2.0
    assert coca_hn["days_in_inventory"] == 45.0

    # P2 tại WH1: xuất 2 trong 90 ngày (daily_rate = 2/90 = 0.022). Tồn 50 -> DSI = 50 / 0.022 = 2250 ngày (>= 90) -> Chậm luân chuyển!
    ruou_hn = sku_items["SKU-RUOU_WH-HN"]
    assert ruou_hn["is_slow_moving"] is True
    assert ruou_hn["days_in_inventory"] >= 90.0

    # P3 tại WH1: xuất 0 trong 90 ngày -> Chậm luân chuyển (không có xuất kho)!
    bap_hn = sku_items["SKU-BAP_WH-HN"]
    assert bap_hn["is_slow_moving"] is True
    assert bap_hn["days_in_inventory"] == 999.0
    assert "Không có xuất kho" in bap_hn["alert"]

    # Thống kê tổng số hàng chậm luân chuyển: 2 mặt hàng (Rượu và Bắp), giá trị = 20tr + 9tr = 29tr
    assert data["slow_moving_skus_count"] == 2
    assert data["slow_moving_value"] == 29000000.0


def test_filter_slow_moving_only(scrum74_env):
    """Lọc chỉ xem danh sách các mặt hàng chậm luân chuyển đang chôn vốn."""
    client, SessionLocal = scrum74_env
    headers = make_auth_header("Warehouse Manager", "wh_mgr")

    res = client.get("/api/v1/reports/inventory-turnover?slow_moving_only=true", headers=headers)
    assert res.status_code == 200
    data = res.json()["data"]
    # Chỉ trả về 2 SKU bị chậm luân chuyển
    assert len(data["sku_details"]) == 2
    for it in data["sku_details"]:
        assert it["is_slow_moving"] is True


def test_filter_by_warehouse_and_category(scrum74_env):
    """Lọc báo cáo theo kho và theo nhóm hàng."""
    client, SessionLocal = scrum74_env
    headers = make_auth_header("Warehouse Manager", "wh_mgr")

    # Lọc theo Kho HCM (WH 2)
    res_wh2 = client.get("/api/v1/reports/inventory-turnover?warehouse_id=2", headers=headers)
    assert res_wh2.status_code == 200
    data_wh2 = res_wh2.json()["data"]
    assert len(data_wh2["sku_details"]) == 1
    assert data_wh2["sku_details"][0]["sku"] == "SKU-COCA"
    assert data_wh2["total_inventory_value"] == 2400000.0

    # Lọc theo Nhóm Thực phẩm khô (Category 2)
    res_cat2 = client.get("/api/v1/reports/inventory-turnover?category_id=2", headers=headers)
    assert res_cat2.status_code == 200
    data_cat2 = res_cat2.json()["data"]
    assert len(data_cat2["sku_details"]) == 1
    assert data_cat2["sku_details"][0]["sku"] == "SKU-BAP"


def test_export_inventory_turnover_excel_success(scrum74_env):
    """AC4: Xuất báo cáo ra file Excel (.xlsx) chuẩn."""
    client, SessionLocal = scrum74_env
    headers = make_auth_header("Warehouse Manager", "wh_mgr")

    res = client.get("/api/v1/reports/inventory-turnover/export?period_days=90&threshold_days=90", headers=headers)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert "attachment; filename=Bao_cao_ton_kho_vong_quay_" in res.headers["content-disposition"]

    # Đọc lại nội dung file Excel bằng openpyxl để xác minh tính hợp lệ
    excel_bytes = io.BytesIO(res.content)
    wb = openpyxl.load_workbook(excel_bytes)

    # Phải có 2 sheet: Chi tiết SKU và Tổng hợp theo Kho & Nhóm
    assert "Chi tiết SKU & Vòng quay" in wb.sheetnames
    assert "Giá trị theo Kho & Nhóm" in wb.sheetnames

    ws1 = wb["Chi tiết SKU & Vòng quay"]
    # Kiểm tra tiêu đề và dòng header
    assert ws1["A1"].value == "BÁO CÁO TỒN KHO VÀ VÒNG QUAY HÀNG HÓA"
    assert ws1["B4"].value == "Mã SKU"
    assert ws1["K4"].value == "Giá trị tồn (VND)"

    ws2 = wb["Giá trị theo Kho & Nhóm"]
    assert ws2["A1"].value == "TỔNG HỢP GIÁ TRỊ TỒN KHO THEO KHO HÀNG"


def test_unauthorized_role_cannot_view_or_export_report(scrum74_env):
    """Phân quyền: Đại lý (Customer) không được phép xem hoặc xuất báo cáo kho."""
    client, SessionLocal = scrum74_env
    cust_headers = make_auth_header("Customer", "cust_user")

    res_view = client.get("/api/v1/reports/inventory-turnover", headers=cust_headers)
    assert res_view.status_code == 403

    res_export = client.get("/api/v1/reports/inventory-turnover/export", headers=cust_headers)
    assert res_export.status_code == 403
