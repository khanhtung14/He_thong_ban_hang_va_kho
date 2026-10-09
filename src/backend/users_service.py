"""Business Service for User Management (SCRUM-62)."""

import csv
import io
import logging
import re
import secrets
import string
from typing import Any, Dict, Iterable, List, Optional, Tuple

import bcrypt
from fastapi import HTTPException, status
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from src.backend.email_service import email_service
from src.backend.models import (
    AccountAuditLog,
    AccountStatus,
    Role,
    Territory,
    User,
    Warehouse,
    seed_default_roles,
)
from src.backend.schemas import UserCreate, UserUpdate

logger = logging.getLogger("users_service")
WAREHOUSE_ROLE_CODES = {"WAREHOUSE", "WH_MANAGER", "WAREHOUSE_KEEPER", "THU_KHO", "WAREHOUSE_STAFF"}


def generate_temporary_password(length: int = 12) -> str:
    """
    Generate a strong temporary password containing:
    - At least 1 uppercase letter
    - At least 1 lowercase letter
    - At least 1 digit
    - At least 1 special symbol
    Minimum length: 8 characters.
    """
    if length < 8:
        length = 8

    uppercase = string.ascii_uppercase
    lowercase = string.ascii_lowercase
    digits = string.digits
    symbols = "!@#$%^&*()_+"

    # Ensure at least one from each category
    password = [
        secrets.choice(uppercase),
        secrets.choice(lowercase),
        secrets.choice(digits),
        secrets.choice(symbols),
    ]

    all_chars = uppercase + lowercase + digits + symbols
    for _ in range(length - 4):
        password.append(secrets.choice(all_chars))

    # Shuffle to eliminate predictable category order
    shuffled = list(password)
    # Using secrets to safely shuffle
    for i in range(len(shuffled) - 1, 0, -1):
        j = secrets.randbelow(i + 1)
        shuffled[i], shuffled[j] = shuffled[j], shuffled[i]

    return "".join(shuffled)


def hash_password(plain_password: str) -> str:
    """Hash a plaintext password using bcrypt with a salt."""
    salt = bcrypt.gensalt()
    return bcrypt.hashpw(plain_password.encode("utf-8"), salt).decode("utf-8")


def verify_password(plain_password: str, hashed_password: str) -> bool:
    """Verify a plaintext password against a bcrypt hash."""
    try:
        return bcrypt.checkpw(plain_password.encode("utf-8"), hashed_password.encode("utf-8"))
    except (ValueError, TypeError):
        return False


def check_duplicate_user(
    db: Session,
    username: Optional[str] = None,
    email: Optional[str] = None,
    phone: Optional[str] = None,
    exclude_user_id: Optional[int] = None,
):
    """
    Validate that username, email, and phone are unique in the database (SCRUM-102).
    Raise an informative 400 Bad Request exception if a collision is detected.
    """
    collision_fields = []

    if username:
        query = db.query(User).filter(func.lower(User.username) == username.strip().lower())
        if exclude_user_id is not None:
            query = query.filter(User.id != exclude_user_id)
        if query.first():
            collision_fields.append("Tên tài khoản")

    if email:
        query = db.query(User).filter(func.lower(User.email) == email.strip().lower())
        if exclude_user_id is not None:
            query = query.filter(User.id != exclude_user_id)
        if query.first():
            collision_fields.append("Email")

    if phone and phone.strip():
        query = db.query(User).filter(User.phone == phone.strip())
        if exclude_user_id is not None:
            query = query.filter(User.id != exclude_user_id)
        if query.first():
            collision_fields.append("Số điện thoại")

    if collision_fields:
        if "Tên tài khoản" in collision_fields and "Email" in collision_fields:
            detail_msg = "Tên tài khoản hoặc Email đã tồn tại trong hệ thống. Vui lòng kiểm tra lại."
        else:
            joined = ", ".join(collision_fields)
            detail_msg = f"{joined} đã tồn tại trong hệ thống. Vui lòng kiểm tra lại."
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=detail_msg)


def resolve_roles(
    db: Session,
    role_codes: Optional[List[str]] = None,
    role: Optional[str] = None
) -> List[Role]:
    """Resolve role codes or names to Role model entities."""
    codes: List[str] = []
    if role_codes:
        codes.extend(role_codes)
    if role and role not in codes:
        codes.append(role)

    # If no role provided, default to SALES_REP per business requirements
    if not codes:
        codes = ["SALES_REP"]

    roles_to_assign: List[Role] = []
    seed_default_roles(db)

    for code in codes:
        clean_code = code.strip()
        if not clean_code:
            continue
        role_entity = (
            db.query(Role)
            .filter(
                or_(
                    func.upper(Role.code) == clean_code.upper(),
                    func.lower(Role.name) == clean_code.lower()
                )
            )
            .first()
        )
        if not role_entity:
            # Map common human readable variants
            variant_map = {
                "SALES": "SALES_REP",
                # Older demo databases store the same role under the code SALES.
                "SALES_REP": "SALES",
                "SALES REP": "SALES_REP",
                "NHAN VIEN KINH DOANH": "SALES_REP",
                "NHÂN VIÊN KINH DOANH": "SALES_REP",
                "ADMIN": "ADMIN",
                "QUẢN TRỊ HỆ THỐNG": "ADMIN",
                "QUAN TRI HE THONG": "ADMIN",
                "SALES MANAGER": "SALES_MANAGER",
                "QUẢN LÝ KINH DOANH": "SALES_MANAGER",
                "WAREHOUSE": "WAREHOUSE",
                "THỦ KHO": "WAREHOUSE",
                "NHÂN VIÊN KHO": "WAREHOUSE",
                "WH MANAGER": "WH_MANAGER",
                "QUẢN LÝ KHO": "WH_MANAGER",
                "ACCOUNTANT": "ACCOUNTANT",
                "KẾ TOÁN CÔNG NỢ": "ACCOUNTANT",
                "CUSTOMER": "CUSTOMER",
                "ĐẠI LÝ": "CUSTOMER",
            }
            mapped_code = variant_map.get(clean_code.upper())
            if mapped_code:
                role_entity = db.query(Role).filter(Role.code == mapped_code).first()

        if not role_entity:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Vai trò '{clean_code}' không hợp lệ hoặc không tồn tại trong hệ thống."
            )
        if role_entity not in roles_to_assign:
            roles_to_assign.append(role_entity)

    return roles_to_assign


def create_user(
    db: Session,
    user_data: UserCreate,
    login_url: str = "/login",
    actor_user_id: Optional[int] = None
) -> Tuple[User, bool, str]:
    """
    Create a new user:
    1. Check duplicates for username, email, and phone (SCRUM-102).
    2. Generate random strong temporary password (SCRUM-100).
    3. Hash password with bcrypt (SCRUM-100).
    4. Set must_change_password=True and initial status (SCRUM-100).
    5. Resolve and associate role(s).
    6. Save User in database.
    7. Record audit log (SCRUM-106).
    8. Dispatch activation email with expiring token & retry (SCRUM-101).
    Returns: (new_user, email_sent, temp_password)
    """
    username = user_data.username.strip()
    full_name = user_data.full_name.strip()
    email = user_data.email.strip().lower()
    phone = user_data.phone.strip() if user_data.phone else None

    # Step 1: Check duplicate
    check_duplicate_user(db, username=username, email=email, phone=phone)

    # Step 2: Generate temporary password
    temp_password = generate_temporary_password(12)

    # Step 3: Hash password
    password_hash = hash_password(temp_password)

    # Step 4: Resolve roles
    roles = resolve_roles(db, role_codes=user_data.role_codes, role=user_data.role)
    selected_warehouse_ids = set(user_data.warehouse_ids)
    warehouses = db.query(Warehouse).filter(
        Warehouse.id.in_(selected_warehouse_ids), Warehouse.is_active.is_(True)
    ).all() if selected_warehouse_ids else []
    if len(warehouses) != len(selected_warehouse_ids):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Có kho được chọn không tồn tại hoặc đã ngừng hoạt động.")
    if any(role.code.upper() in WAREHOUSE_ROLE_CODES for role in roles) and not warehouses and db.query(Warehouse).filter(Warehouse.is_active.is_(True)).first() is not None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Người dùng thuộc vai trò kho phải được gắn với ít nhất một kho.")
    selected_territory_ids = set(user_data.territory_ids)
    territories = db.query(Territory).filter(Territory.id.in_(selected_territory_ids)).all() if selected_territory_ids else []
    if len(territories) != len(selected_territory_ids):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Có địa bàn được chọn không tồn tại.")

    # Step 5: Save User to DB with must_change_password=True & status
    initial_status = user_data.status or AccountStatus.PENDING_ACTIVATION
    new_user = User(
        username=username,
        email=email,
        full_name=full_name,
        phone=phone,
        password_hash=password_hash,
        status=initial_status,
        must_change_password=True,
        roles=roles,
        warehouses=warehouses,
        territories=territories,
    )

    db.add(new_user)
    try:
        db.commit()
        db.refresh(new_user)
    except IntegrityError as exc:
        db.rollback()
        logger.error(f"Database IntegrityError while creating user: {exc}")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Tên tài khoản hoặc Email đã tồn tại trong hệ thống. Vui lòng kiểm tra lại."
        ) from exc

    # Step 6: Record Audit Log (SCRUM-106)
    try:
        audit_log = AccountAuditLog(
            user_id=new_user.id,
            actor_user_id=actor_user_id,
            action="CREATE_USER",
            reason=f"Khởi tạo tài khoản {new_user.username} với vai trò {[r.code for r in roles]}",
        )
        db.add(audit_log)
        db.commit()
    except Exception as audit_err:
        logger.warning(f"Failed to write audit log for user creation: {audit_err}")
        db.rollback()

    # Step 7: Dispatch activation email with expiration token and retry mechanism (SCRUM-101)
    activation_token = secrets.token_urlsafe(32)
    email_sent = email_service.send_activation_email(
        to_email=new_user.email,
        full_name=new_user.full_name,
        username=new_user.username,
        temp_password=temp_password,
        login_url=login_url,
        activation_token=activation_token,
        max_retries=3
    )

    return new_user, email_sent, temp_password


def normalize_import_header(value: Optional[str]) -> str:
    if value is None:
        return ""
    cleaned = value.strip().lower().replace("-", " ").replace("_", " ")
    cleaned = " ".join(cleaned.split())
    aliases = {
        "ho ten": "full_name",
        "họ tên": "full_name",
        "họ và tên": "full_name",
        "ho va ten": "full_name",
        "full name": "full_name",
        "fullname": "full_name",
        "ten nguoi dung": "username",
        "user name": "username",
        "ten dang nhap": "username",
        "tên đăng nhập": "username",
        "tài khoản": "username",
        "email": "email",
        "địa chỉ email": "email",
        "phone": "phone",
        "dien thoai": "phone",
        "số điện thoại": "phone",
        "sdt": "phone",
        "mobile": "phone",
        "role code": "role_codes",
        "role codes": "role_codes",
        "vai tro": "role_codes",
        "vai trò": "role_codes",
        "mã vai trò": "role_codes",
        "roles": "role_codes",
        "role": "role_codes",
    }
    return aliases.get(cleaned, cleaned.replace(" ", "_"))


def parse_excel_rows(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Parse an Excel file or CSV into a list of normalized dictionaries."""
    required_headers = {"full_name", "username", "email"}
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        headers = {normalize_import_header(key) for key in (reader.fieldnames or []) if key}
        missing_headers = required_headers - headers
        if missing_headers:
            raise ValueError("Thiếu cột bắt buộc trong tiêu đề: Họ tên, Tên đăng nhập, Email.")
        results: List[Dict[str, Any]] = []
        for row in reader:
            if not row or not any(str(value or "").strip() for value in row.values()):
                continue
            mapping = {
                normalize_import_header(key): (value.strip() if isinstance(value, str) else value)
                for key, value in row.items() if key is not None
            }
            mapping["__import_row_number"] = reader.line_num
            results.append(mapping)
        return results

    workbook = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header_index = None
    header: List[str] = []
    for candidate_index, candidate in enumerate(rows[:10]):
        candidate_headers = [normalize_import_header(str(cell).strip()) if cell is not None else "" for cell in candidate]
        if required_headers.issubset(set(candidate_headers)):
            header_index = candidate_index
            header = candidate_headers
            break
    if header_index is None:
        raise ValueError("Không tìm thấy hàng tiêu đề có đủ cột Họ tên, Tên đăng nhập và Email trong 10 hàng đầu.")

    results: List[Dict[str, Any]] = []
    for source_row_number, values in enumerate(rows[header_index + 1:], start=header_index + 2):
        mapping: Dict[str, Any] = {}
        for key, value in zip(header, values):
            if not key:
                continue
            mapping[key] = value.strip() if isinstance(value, str) else value
        if mapping and any((str(value)).strip() for value in mapping.values() if value is not None):
            mapping["__import_row_number"] = source_row_number
            results.append(mapping)
    return results


def normalize_import_phone(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None

    digits = re.sub(r"\D", "", text)
    if len(digits) == 11 and digits.startswith("84"):
        digits = "0" + digits[2:]
    if len(digits) == 10 and digits.startswith("0"):
        return digits
    return text


def bulk_create_users_from_rows(db: Session, rows: Iterable[Dict[str, Any]], actor_user_id: Optional[int] = None) -> Tuple[int, List[str], int]:
    """Create multiple users from imported rows and skip duplicates silently."""
    created_count = 0
    errors: List[str] = []
    duplicate_count = 0

    for row_index, row in enumerate(rows, start=2):
        if row is None:
            continue
        source_row_index = int(row.get("__import_row_number", row_index) or row_index)
        normalized = {normalize_import_header(str(key)): value for key, value in row.items() if key is not None}
        if not normalized:
            continue

        full_name = str(normalized.get("full_name") or normalized.get("ho_ten") or normalized.get("name") or "").strip()
        username = str(normalized.get("username") or normalized.get("ten_dang_nhap") or normalized.get("user_name") or "").strip()
        email = str(normalized.get("email") or "").strip().lower()
        phone = normalize_import_phone(normalized.get("phone") or normalized.get("dien_thoai") or normalized.get("sdt") or "")
        raw_role_codes = normalized.get("role_codes") or normalized.get("role") or normalized.get("vai_tro") or "SALES_REP"

        header_tokens = {
            "full_name", "ho_ten", "name", "username", "user_name", "ten_dang_nhap",
            "email", "phone", "dien_thoai", "sdt", "role_codes", "role", "vai_tro"
        }
        if full_name.lower() in header_tokens or username.lower() in header_tokens or email.lower() in header_tokens:
            continue
        if not full_name and not username and not email and not phone:
            continue

        if not full_name or not username or not email:
            errors.append(f"Dòng {source_row_index}: thiếu họ tên, tên đăng nhập hoặc email.")
            continue

        parsed_roles = []
        if isinstance(raw_role_codes, str):
            parsed_roles = [part.strip() for part in raw_role_codes.replace(";", ",").split(",") if part.strip()]
        elif isinstance(raw_role_codes, (list, tuple)):
            parsed_roles = [str(part).strip() for part in raw_role_codes if str(part).strip()]
        if not parsed_roles:
            parsed_roles = ["SALES_REP"]

        try:
            model = UserCreate(
                username=username,
                full_name=full_name,
                email=email,
                phone=phone,
                role_codes=parsed_roles,
                status=AccountStatus.PENDING_ACTIVATION,
            )
            create_user(db=db, user_data=model, actor_user_id=actor_user_id)
            created_count += 1
        except HTTPException as exc:
            detail = str(exc.detail or "")
            if "đã tồn tại" in detail.lower() or "already" in detail.lower() or "tồn tại" in detail.lower():
                duplicate_count += 1
                continue
            errors.append(f"Dòng {source_row_index}: {detail}")
        except Exception as exc:  # pragma: no cover - defensive fallback
            errors.append(f"Dòng {source_row_index}: {exc}")

    return created_count, errors, duplicate_count


def preview_user_import_rows(db: Session, rows: Iterable[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Validate imported user rows without creating user records."""
    preview: List[Dict[str, Any]] = []
    seen: Dict[str, int] = {}
    header_tokens = {
        "full_name", "ho_ten", "name", "username", "user_name", "ten_dang_nhap",
        "email", "phone", "dien_thoai", "sdt", "role_codes", "role", "vai_tro"
    }

    for row_index, row in enumerate(rows, start=2):
        normalized = {normalize_import_header(str(key)): value for key, value in row.items() if key is not None}
        source_row_index = int(normalized.get("__import_row_number", row_index) or row_index)
        full_name = str(normalized.get("full_name") or normalized.get("ho_ten") or normalized.get("name") or "").strip()
        username = str(normalized.get("username") or normalized.get("ten_dang_nhap") or normalized.get("user_name") or "").strip()
        email = str(normalized.get("email") or "").strip().lower()
        phone = normalize_import_phone(normalized.get("phone") or normalized.get("dien_thoai") or normalized.get("sdt") or "")
        raw_role_codes = normalized.get("role_codes") or normalized.get("role") or normalized.get("vai_tro") or "SALES_REP"
        if not normalized or full_name.lower() in header_tokens or username.lower() in header_tokens or email.lower() in header_tokens:
            continue
        if not full_name and not username and not email and not phone:
            continue

        if isinstance(raw_role_codes, str):
            role_codes = [part.strip() for part in raw_role_codes.replace(";", ",").split(",") if part.strip()]
        elif isinstance(raw_role_codes, (list, tuple)):
            role_codes = [str(part).strip() for part in raw_role_codes if str(part).strip()]
        else:
            role_codes = []
        if not role_codes:
            role_codes = ["SALES_REP"]

        error = ""
        duplicate = False
        try:
            user_data = UserCreate(
                username=username,
                full_name=full_name,
                email=email,
                phone=phone,
                role_codes=role_codes,
                status=AccountStatus.PENDING_ACTIVATION,
            )
            check_duplicate_user(db, username=username, email=email, phone=phone)
            roles = resolve_roles(db, role_codes=user_data.role_codes, role=user_data.role)
            warehouse_is_required = any(role.code.upper() in WAREHOUSE_ROLE_CODES for role in roles)
            active_warehouse_exists = db.query(Warehouse).filter(Warehouse.is_active.is_(True)).first() is not None
            if warehouse_is_required and active_warehouse_exists:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Người dùng thuộc vai trò kho phải được gắn với ít nhất một kho.",
                )
            row_keys = [f"username:{username.lower()}", f"email:{email.lower()}"]
            if phone:
                row_keys.append(f"phone:{phone}")
            duplicate_line = next((seen[key] for key in row_keys if key in seen), None)
            if duplicate_line is not None:
                error = f"Trùng thông tin với dòng {duplicate_line}."
                duplicate = True
            else:
                for key in row_keys:
                    seen[key] = source_row_index
        except HTTPException as exc:
            error = str(exc.detail or "Dữ liệu không hợp lệ.")
            duplicate = "tồn tại" in error.lower() or "đã tồn tại" in error.lower()
        except ValidationError as exc:
            field_labels = {"username": "Tên đăng nhập", "full_name": "Họ tên", "email": "Email", "phone": "Số điện thoại"}
            messages = []
            for item in exc.errors():
                field = str(item["loc"][0]) if item.get("loc") else "dữ liệu"
                label = field_labels.get(field, field)
                if item.get("type") == "missing":
                    messages.append(f"Thiếu {label.lower()}.")
                elif field == "email":
                    messages.append("Email không hợp lệ.")
                elif item.get("type") == "string_too_short" and field == "username":
                    messages.append("Tên đăng nhập phải có ít nhất 3 ký tự.")
                elif item.get("type") == "string_too_short" and field == "full_name":
                    messages.append("Họ tên không được để trống.")
                else:
                    messages.append(str(item.get("msg", "Dữ liệu không hợp lệ.")).replace("Value error, ", ""))
            error = " ".join(messages)
        except Exception as exc:
            error = str(exc)

        preview.append({
            "row": source_row_index,
            "full_name": full_name,
            "username": username,
            "email": email,
            "phone": phone,
            "role_codes": role_codes,
            "valid": not error,
            "duplicate": duplicate,
            "error": error,
        })

    return preview


def get_user_by_id(db: Session, user_id: int) -> User:
    """Retrieve user by ID or raise HTTP 404."""
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Không tìm thấy người dùng có ID {user_id}."
        )
    return user


def update_user(
    db: Session,
    user_id: int,
    user_data: UserUpdate,
    actor_user_id: Optional[int] = None
) -> User:
    """
    Update user information: full_name, email, phone, role(s), status (SCRUM-103).
    Record audit log for changes including lock/unlock operations (SCRUM-106).
    """
    user = get_user_by_id(db, user_id)
    old_status = user.status

    new_email = user_data.email.strip().lower() if user_data.email else None
    new_phone = user_data.phone.strip() if user_data.phone is not None else None

    # Check duplicates if email or phone is updated
    check_duplicate_user(
        db,
        email=new_email if (new_email and new_email != user.email) else None,
        phone=new_phone if (new_phone and new_phone != user.phone) else None,
        exclude_user_id=user_id
    )

    updated_fields = []

    if user_data.full_name is not None:
        user.full_name = user_data.full_name.strip()
        updated_fields.append("full_name")

    if new_email is not None:
        user.email = new_email
        updated_fields.append("email")

    if user_data.phone is not None:
        user.phone = new_phone
        updated_fields.append("phone")

    if user_data.status is not None:
        user.status = user_data.status
        updated_fields.append(f"status:{user_data.status.value}")

    if user_data.must_change_password is not None:
        user.must_change_password = user_data.must_change_password
        updated_fields.append("must_change_password")

    if user_data.role_codes is not None or user_data.role is not None:
        roles = resolve_roles(db, role_codes=user_data.role_codes, role=user_data.role)
        current_admin = any(role.code.upper() == "ADMIN" for role in user.roles)
        assigned_admin = any(role.code.upper() == "ADMIN" for role in roles)
        if actor_user_id == user.id and current_admin and not assigned_admin:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Không thể tự thu hồi vai trò quản trị của chính mình.",
            )
        if any(role.code.upper() in WAREHOUSE_ROLE_CODES for role in roles) and not user.warehouses and db.query(Warehouse).filter(Warehouse.is_active.is_(True)).first() is not None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Người dùng thuộc vai trò kho phải được gắn với ít nhất một kho.",
            )
        user.roles = roles
        updated_fields.append(f"roles:{[r.code for r in roles]}")

    try:
        db.commit()
        db.refresh(user)
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Thông tin cập nhật bị trùng lặp với người dùng khác."
        ) from exc

    # Determine action for audit log (SCRUM-106)
    action = "UPDATE_USER"
    if user_data.status == AccountStatus.LOCKED and old_status != AccountStatus.LOCKED:
        action = "LOCK_USER"
    elif user_data.status == AccountStatus.ACTIVE and old_status == AccountStatus.LOCKED:
        action = "UNLOCK_USER"

    try:
        audit_log = AccountAuditLog(
            user_id=user.id,
            actor_user_id=actor_user_id,
            action=action,
            reason=f"Cập nhật tài khoản {user.username}: {', '.join(updated_fields)}",
        )
        db.add(audit_log)
        db.commit()
    except Exception as audit_err:
        logger.warning(f"Failed to write audit log for user update: {audit_err}")
        db.rollback()
    return user


def list_users(
    db: Session,
    search: Optional[str] = None,
    role: Optional[str] = None,
    status_filter: Optional[AccountStatus] = None,
    page: int = 1,
    page_size: int = 20
) -> Tuple[List[User], int, int]:
    """
    Search, filter, and paginate users at the database query level.
    Default page_size = 20.
    Returns: (users, total_count, total_pages)
    """
    page = max(1, page)
    page_size = max(1, min(page_size, 100))

    query = db.query(User)

    # Search by full_name, username, or phone
    if search and search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                User.full_name.ilike(term),
                User.username.ilike(term),
                User.phone.ilike(term)
            )
        )

    # Filter by role
    if role and role.strip():
        role_term = role.strip()
        query = query.join(User.roles).filter(
            or_(
                func.upper(Role.code) == role_term.upper(),
                Role.name.ilike(f"%{role_term}%")
            )
        )

    # Filter by status
    if status_filter:
        query = query.filter(User.status == status_filter)

    total_count = query.distinct().count()
    total_pages = (total_count + page_size - 1) // page_size if total_count > 0 else 1

    users = (
        query.distinct()
        .options(selectinload(User.roles), selectinload(User.territories))
        .order_by(User.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )

    return users, total_count, total_pages
