from io import BytesIO
import re

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook


router = APIRouter()

# Dữ liệu tạm để kiểm thử.
# Khi tích hợp database thật, thay danh sách này bằng truy vấn DB.
imported_users = []


REQUIRED_HEADERS = [
    "username",
    "full_name",
    "email",
    "role",
]


def create_excel_template():
    """Tạo file Excel mẫu để quản trị viên tải về."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sach nguoi dung"

    sheet.append(["username", "full_name", "email", "role"])
    sheet.append([
        "nguyenvana",
        "Nguyen Van A",
        "vana@example.com",
        "sales",
    ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def read_excel_file(file_content):
    """Đọc dữ liệu từ file Excel và kiểm tra tiêu đề cột."""
    try:
        workbook = load_workbook(
            BytesIO(file_content),
            read_only=True,
            data_only=True,
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Không đọc được file Excel. Vui lòng kiểm tra lại file.",
        )

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    headers = next(rows, None)

    if not headers:
        raise HTTPException(
            status_code=400,
            detail="File Excel đang trống.",
        )

    headers = [
        str(value).strip().lower() if value is not None else ""
        for value in headers
    ]

    missing_headers = [
        header
        for header in REQUIRED_HEADERS
        if header not in headers
    ]

    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "File Excel thiếu cột bắt buộc.",
                "missing_headers": missing_headers,
            },
        )

    header_positions = {
        header: headers.index(header)
        for header in REQUIRED_HEADERS
    }

    result = []

    for row_number, row in enumerate(rows, start=2):
        user_data = {}

        for header in REQUIRED_HEADERS:
            position = header_positions[header]
            value = row[position] if position < len(row) else None

            user_data[header] = (
                str(value).strip() if value is not None else ""
            )

        # Bỏ qua dòng hoàn toàn trống.
        if not any(user_data.values()):
            continue

        result.append({
            "row_number": row_number,
            "data": user_data,
        })

    workbook.close()
    return result


def validate_rows(rows):
    """Kiểm tra dữ liệu từng dòng, không lưu dữ liệu."""
    results = []

    seen_usernames = set()
    seen_emails = set()

    existing_usernames = {
        user["username"].lower()
        for user in imported_users
    }
    existing_emails = {
        user["email"].lower()
        for user in imported_users
    }

    for item in rows:
        row_number = item["row_number"]
        data = item["data"]
        errors = []

        username = data["username"]
        full_name = data["full_name"]
        email = data["email"]
        role = data["role"]

        if not username:
            errors.append("Thiếu tên đăng nhập.")

        if not full_name:
            errors.append("Thiếu họ và tên.")

        if not email:
            errors.append("Thiếu email.")
        elif not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            errors.append("Email không đúng định dạng.")

        if not role:
            errors.append("Thiếu vai trò.")

        if username:
            username_key = username.lower()

            if username_key in existing_usernames:
                errors.append("Tên đăng nhập đã tồn tại.")

            if username_key in seen_usernames:
                errors.append("Tên đăng nhập bị trùng trong file.")

            seen_usernames.add(username_key)

        if email:
            email_key = email.lower()

            if email_key in existing_emails:
                errors.append("Email đã tồn tại.")

            if email_key in seen_emails:
                errors.append("Email bị trùng trong file.")

            seen_emails.add(email_key)

        results.append({
            "row_number": row_number,
            "data": data,
            "valid": len(errors) == 0,
            "errors": errors,
        })

    return results


@router.get("/users/import/template")
def download_users_template():
    """Tải file Excel mẫu."""
    excel_file = create_excel_template()

    return StreamingResponse(
        excel_file,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                "attachment; filename=users_template.xlsx"
            )
        },
    )


@router.post("/users/import/preview")
async def preview_users_import(file: UploadFile = File(...)):
    """Xem trước dữ liệu và báo lỗi theo từng dòng."""
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Vui lòng tải lên file Excel có đuôi .xlsx.",
        )

    content = await file.read()
    rows = read_excel_file(content)
    results = validate_rows(rows)

    valid_count = sum(1 for item in results if item["valid"])
    error_count = len(results) - valid_count

    return {
        "total_rows": len(results),
        "valid_rows": valid_count,
        "error_rows": error_count,
        "rows": results,
    }


@router.post("/users/import")
async def import_users(file: UploadFile = File(...)):
    """Nhập các dòng hợp lệ và bỏ qua các dòng lỗi."""
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Vui lòng tải lên file Excel có đuôi .xlsx.",
        )

    content = await file.read()
    rows = read_excel_file(content)
    results = validate_rows(rows)

    created_count = 0
    skipped_count = 0
    report = []

    for item in results:
        if not item["valid"]:
            skipped_count += 1
            report.append({
                "row_number": item["row_number"],
                "status": "skipped",
                "errors": item["errors"],
            })
            continue

        user_data = item["data"]

        imported_users.append({
            "username": user_data["username"],
            "full_name": user_data["full_name"],
            "email": user_data["email"],
            "role": user_data["role"],
        })

        created_count += 1
        report.append({
            "row_number": item["row_number"],
            "status": "created",
            "errors": [],
        })

    return {
        "total_rows": len(results),
        "created_count": created_count,
        "skipped_count": skipped_count,
        "report": report,
    }