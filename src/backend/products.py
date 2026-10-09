"""Products endpoints with RBAC enforcement and sensitive financial data masking."""

import io
import math
from typing import Any, Dict
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, HTTPException, Path as PathParam, UploadFile, status
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, Field, field_validator
from sqlalchemy.orm import Session
from starlette.responses import StreamingResponse

try:
    from src.backend.rbac import (
        PERM_PRODUCTS_IMPORT,
        PERM_PRODUCTS_VIEW,
        PERM_PRODUCTS_UNIT_MANAGE,
        AuthenticatedUser,
        require_permissions,
        sanitize_financial_data,
    )
    from src.backend.database import get_db
    from src.backend.models import ProductUnit
except ModuleNotFoundError:  # pragma: no cover
    from rbac import (
        PERM_PRODUCTS_IMPORT,
        PERM_PRODUCTS_VIEW,
        PERM_PRODUCTS_UNIT_MANAGE,
        AuthenticatedUser,
        require_permissions,
        sanitize_financial_data,
    )
    from database import get_db
    from models import ProductUnit

router = APIRouter(prefix="/api/v1/products", tags=["Products"])
MAX_IMPORT_BYTES = 10 * 1024 * 1024
IMPORT_COLUMNS = (
    "sku",
    "name",
    "category",
    "sale_price",
    "cost_price",
    "stock_available",
    "unit",
)
REQUIRED_IMPORT_COLUMNS = {"sku", "name", "category"}

# Mock products database
MOCK_PRODUCTS: Dict[str, Dict[str, Any]] = {
    "SKU-001": {
        "sku": "SKU-001",
        "name": "Nước tăng lực Red Bull 250ml",
        "category": "Nước giải khát",
        "sale_price": 500000,
        "cost_price": 350000,
        "margin": "30.0%",
        "stock_available": 120,
        "unit": "Thùng 24 lon",
    },
    "SKU-002": {
        "sku": "SKU-002",
        "name": "Cà phê lon Highlands 235ml",
        "category": "Cà phê",
        "sale_price": 240000,
        "cost_price": 180000,
        "margin": "25.0%",
        "stock_available": 85,
        "unit": "Thùng 24 lon",
    },
}


class ProductUnitRequest(BaseModel):
    unit_name: str = Field(min_length=1, max_length=100)
    conversion_rate: float = Field(gt=0, allow_inf_nan=False)
    is_base_unit: bool = False

    @field_validator("unit_name")
    @classmethod
    def clean_unit_name(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("Tên đơn vị không được để trống.")
        return value


@router.get("")
def list_products(
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Get list of products.

    Confidentiality rule:
    - Warehouse and Sales Rep: cost_price and margin are stripped.
    - Sales Manager only: cost_price and margin are present.
    """
    products_list = [p.copy() for p in MOCK_PRODUCTS.values()]
    return sanitize_financial_data(products_list, user)


@router.get("/{sku}")
def get_product_detail(
    sku: str,
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
) -> Any:
    """Get product detail by SKU.

    Confidentiality rule:
    - Warehouse and Sales Rep: cost_price and margin are stripped.
    - Sales Manager only: cost_price and margin are present.
    """
    product = MOCK_PRODUCTS.get(sku.upper()) or MOCK_PRODUCTS.get(sku)
    if not product:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Sản phẩm '{sku}' không tồn tại.")

    sanitized = sanitize_financial_data(product.copy(), user)
    return sanitized


@router.get("/{product_id}/units")
def list_product_units(
    product_id: int = PathParam(gt=0),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_VIEW)),
    db: Session = Depends(get_db),
) -> list[dict[str, Any]]:
    """List configured conversion units for one product."""
    rows = (
        db.query(ProductUnit)
        .filter(ProductUnit.product_id == product_id)
        .order_by(ProductUnit.is_base_unit.desc(), ProductUnit.unit_name.asc())
        .all()
    )
    return [
        {
            "id": row.id,
            "product_id": row.product_id,
            "unit_name": row.unit_name,
            "conversion_rate": row.conversion_rate,
            "is_base_unit": row.is_base_unit,
        }
        for row in rows
    ]


@router.post("/{product_id}/units")
def upsert_product_unit(
    payload: ProductUnitRequest,
    product_id: int = PathParam(gt=0),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_UNIT_MANAGE)),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Create or update a unit; at most one unit per product can be the base."""
    conversion_rate = 1.0 if payload.is_base_unit else payload.conversion_rate
    try:
        row = (
            db.query(ProductUnit)
            .filter(
                ProductUnit.product_id == product_id,
                ProductUnit.unit_name == payload.unit_name,
            )
            .first()
        )
        if row is None:
            row = ProductUnit(
                product_id=product_id,
                unit_name=payload.unit_name,
                conversion_rate=conversion_rate,
                is_base_unit=payload.is_base_unit,
            )
            db.add(row)
        else:
            row.conversion_rate = conversion_rate
            row.is_base_unit = payload.is_base_unit

        if payload.is_base_unit:
            (
                db.query(ProductUnit)
                .filter(
                    ProductUnit.product_id == product_id,
                    ProductUnit.unit_name != payload.unit_name,
                )
                .update({ProductUnit.is_base_unit: False}, synchronize_session=False)
            )
        db.commit()
        db.refresh(row)
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không thể lưu đơn vị tính sản phẩm.",
        ) from exc

    return {
        "id": row.id,
        "product_id": row.product_id,
        "unit_name": row.unit_name,
        "conversion_rate": row.conversion_rate,
        "is_base_unit": row.is_base_unit,
    }


def _row_to_import_item(row_number: int, values: dict[str, Any]) -> dict[str, Any]:
    errors: list[str] = []
    sku = str(values.get("sku") or "").strip().upper()
    name = str(values.get("name") or "").strip()
    category = str(values.get("category") or "").strip()

    if not sku:
        errors.append("SKU không được để trống.")
    if not name:
        errors.append("Tên sản phẩm không được để trống.")
    if not category:
        errors.append("Danh mục không được để trống.")

    item: dict[str, Any] = {"sku": sku, "name": name, "category": category}
    for column in ("sale_price", "cost_price", "stock_available"):
        raw_value = values.get(column)
        if raw_value is None or str(raw_value).strip() == "":
            continue
        try:
            number = float(str(raw_value).strip())
            if not math.isfinite(number) or number < 0 or (
                column == "stock_available" and not number.is_integer()
            ):
                raise ValueError
            item[column] = int(number) if number.is_integer() else number
        except (TypeError, ValueError):
            suffix = " nguyên." if column == "stock_available" else "."
            errors.append(f"{column} phải là số không âm{suffix}")

    unit = str(values.get("unit") or "").strip()
    if unit:
        item["unit"] = unit
    return {"row": row_number, "sku": sku, "item": item, "errors": errors}


def _parse_import_workbook(content: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="Tệp không phải Excel .xlsx hợp lệ.") from exc

    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise HTTPException(status_code=422, detail="Tệp Excel không có dòng tiêu đề.")

        normalized_headers = [str(value or "").strip().lower() for value in headers]
        nonempty_headers = [header for header in normalized_headers if header]
        if len(nonempty_headers) != len(set(nonempty_headers)):
            raise HTTPException(status_code=422, detail="Tệp Excel có tên cột bị trùng.")
        missing_columns = REQUIRED_IMPORT_COLUMNS - set(nonempty_headers)
        if missing_columns:
            raise HTTPException(
                status_code=422,
                detail=f"Thiếu cột bắt buộc: {', '.join(sorted(missing_columns))}.",
            )

        parsed_rows: list[dict[str, Any]] = []
        for row_number, cells in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in cells):
                continue
            values = {
                header: cells[index] if index < len(cells) else None
                for index, header in enumerate(normalized_headers)
                if header in IMPORT_COLUMNS
            }
            parsed_rows.append(_row_to_import_item(row_number, values))
        if not parsed_rows:
            raise HTTPException(status_code=422, detail="Tệp Excel không có dữ liệu sản phẩm.")
        return parsed_rows
    finally:
        workbook.close()


async def _read_import_file(file: UploadFile) -> bytes:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="Chỉ hỗ trợ tệp Excel .xlsx.")
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Tệp Excel vượt quá giới hạn 10 MB.")
    return content


def _annotate_import_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen_skus: set[str] = set()
    existing_skus = {sku.upper() for sku in MOCK_PRODUCTS}
    for row in rows:
        sku = row["sku"]
        if sku and sku in seen_skus:
            row["errors"].append("SKU bị lặp trong tệp Excel.")
        if sku:
            seen_skus.add(sku)
        row["action"] = "error" if row["errors"] else ("update" if sku in existing_skus else "create")
        row["message"] = (
            "; ".join(row["errors"])
            if row["errors"]
            else ("SKU đã tồn tại, dữ liệu sẽ được cập nhật." if row["action"] == "update" else "Sản phẩm sẽ được tạo mới.")
        )
        row.pop("errors", None)
    return rows


@router.get("/import/template")
def download_import_template(
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_IMPORT)),
) -> StreamingResponse:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Products"
    worksheet.append(IMPORT_COLUMNS)
    worksheet.append(["SKU-003", "Tên sản phẩm", "Đồ uống", 100000, 70000, 20, "Thùng"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="product-import-template.xlsx"'},
    )


@router.post("/import/preview")
async def preview_product_import(
    file: UploadFile = File(...),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_IMPORT)),
) -> dict[str, Any]:
    rows = _annotate_import_rows(_parse_import_workbook(await _read_import_file(file)))
    return {
        "total": len(rows),
        "create_count": sum(row["action"] == "create" for row in rows),
        "update_count": sum(row["action"] == "update" for row in rows),
        "error_count": sum(row["action"] == "error" for row in rows),
        "rows": rows,
    }


@router.post("/import/commit")
async def commit_product_import(
    file: UploadFile = File(...),
    user: AuthenticatedUser = Depends(require_permissions(PERM_PRODUCTS_IMPORT)),
) -> dict[str, Any]:
    rows = _annotate_import_rows(_parse_import_workbook(await _read_import_file(file)))
    if any(row["action"] == "error" for row in rows):
        raise HTTPException(
            status_code=422,
            detail={"message": "Không có dữ liệu nào được lưu vì tệp còn dòng lỗi.", "rows": rows},
        )

    for row in rows:
        item = row["item"]
        existing_key = next((key for key in MOCK_PRODUCTS if key.upper() == row["sku"]), None)
        if existing_key is not None:
            MOCK_PRODUCTS[existing_key].update(item)
        else:
            MOCK_PRODUCTS[row["sku"]] = item
    return {
        "message": "Nhập sản phẩm thành công.",
        "created_count": sum(row["action"] == "create" for row in rows),
        "updated_count": sum(row["action"] == "update" for row in rows),
    }
